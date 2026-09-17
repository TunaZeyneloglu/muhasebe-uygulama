import os
import glob
import xml.etree.ElementTree as ET
import pandas as pd
from datetime import datetime
import subprocess
import sys
import unicodedata
from pathlib import Path
import customtkinter as ctk
from tkinter import filedialog
from openpyxl.styles import Font
import warnings

# openpyxl uyarılarını sustur
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

# Tema ayarı
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

MASAUSTU = str(Path.home() / "OneDrive" / "Desktop")
if not os.path.exists(MASAUSTU):
    MASAUSTU = str(Path.home() / "Desktop")

SON_DOSYA_YOLU = None
TARGET_VKN_FOR_ZERO_VAT = "7740042326"

# XML Namespace sabitleri
XML_NS = {
    "cbc": "urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2",
    "cac": "urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2",
}

# ----------------- Ortak Fonksiyonlar -----------------

def cari_adi_al(party, ns):
    cariadi = party.findtext("cac:PartyName/cbc:Name", namespaces=ns)
    if not cariadi:
        firstname = party.findtext("cac:Person/cbc:FirstName", namespaces=ns)
        lastname = party.findtext("cac:Person/cbc:FamilyName", namespaces=ns)
        if firstname and lastname:
            cariadi = f"{firstname} {lastname}"
    if not cariadi:
        cariadi = party.findtext("cac:PartyLegalEntity/cbc:RegistrationName", namespaces=ns)
    return cariadi if cariadi else "Bilinmeyen Cari"

def kaydet_excel(veriler, dosya_yolu):
    df = pd.DataFrame(veriler)[[
        "SIRANO", "TARIH", "EVRAKNO", "EVRAKTURU", "CARIKODU", "CARIADI",
        "STOKKODU", "STOKADI", "BIRIM", "MIKTAR", "DOVIZC",
        "BIRIMFIYAT", "INDTL", "KDVY", "TUTARTL", "FATURACK"
    ]]
    # Metin sütunları: 0 ile başlayan kodların korunması için
    df["EVRAKNO"] = df["EVRAKNO"].fillna("").astype(str)
    df["CARIKODU"] = df["CARIKODU"].fillna("").astype(str)
    df["STOKKODU"] = df["STOKKODU"].fillna("").astype(str)
    with pd.ExcelWriter(dosya_yolu, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Sayfa1")

def dosyayi_ac():
    global SON_DOSYA_YOLU
    if SON_DOSYA_YOLU:
        try:
            if sys.platform == "win32":
                os.startfile(SON_DOSYA_YOLU)
            elif sys.platform == "darwin":
                subprocess.call(["open", SON_DOSYA_YOLU])
            else:
                subprocess.call(["xdg-open", SON_DOSYA_YOLU])
        except Exception:
            pass

# ----------------- Ortak Yardımcı Fonksiyonlar (Karşılaştırma) -----------------

def _temizle_fatura_no(fatura_no):
    """Fatura numarasını temizle ve normalize et"""
    if pd.isna(fatura_no):
        return ""
    return str(fatura_no).strip().replace("'", "").upper()

def _excel_dosyalari_oku(zirve_path, portal_paths):
    """Zirve ve portal Excel dosyalarını oku ve birleştir"""
    df_zirve = pd.read_excel(zirve_path)
    
    if isinstance(portal_paths, list):
        portal_dfs = [pd.read_excel(path) for path in portal_paths]
        df_portal = pd.concat(portal_dfs, ignore_index=True)
    else:
        df_portal = pd.read_excel(portal_paths)
    
    return df_zirve, df_portal

def _kontrol_sonuc_excel_yaz(sonuc_dosyasi, tutar_farklari, eksik_girisler, fazla_girisler):
    """Kontrol sonuçlarını Excel'e yaz"""
    with pd.ExcelWriter(sonuc_dosyasi, engine='openpyxl') as writer:
        # Sayfa 1: Tutar Farkları
        if tutar_farklari:
            df_tutar = pd.DataFrame(tutar_farklari).sort_values('Fatura No')
            df_tutar.to_excel(writer, sheet_name='Tutar Farkları', index=False)
            
            # Formatlama
            worksheet = writer.sheets['Tutar Farkları']
            headers = list(df_tutar.columns)
            tutar_farki_col = headers.index('Tutar Farkı') + 1
            kdv_farki_col = headers.index('KDV Farkı') + 1
            
            for row_idx in range(2, len(df_tutar) + 2):
                worksheet.cell(row=row_idx, column=tutar_farki_col).font = Font(bold=True, size=13)
                worksheet.cell(row=row_idx, column=kdv_farki_col).font = Font(bold=True, size=13)
        else:
            pd.DataFrame({'Mesaj': ['Tutar farkı bulunamadı.']}).to_excel(
                writer, sheet_name='Tutar Farkları', index=False)
        
        # Sayfa 2: Eksik Girişler
        if eksik_girisler:
            pd.DataFrame(eksik_girisler).sort_values('Fatura No').to_excel(
                writer, sheet_name='Eksik Girişler', index=False)
        else:
            pd.DataFrame({'Mesaj': ['Eksik giriş bulunamadı.']}).to_excel(
                writer, sheet_name='Eksik Girişler', index=False)
        
        # Sayfa 3: Fazla Girişler
        if fazla_girisler:
            pd.DataFrame(fazla_girisler).sort_values('Fatura No').to_excel(
                writer, sheet_name='Fazla Girişler', index=False)
        else:
            pd.DataFrame({'Mesaj': ['Fazla giriş bulunamadı.']}).to_excel(
                writer, sheet_name='Fazla Girişler', index=False)

def _kontrol_ozet_mesaj_olustur(portal_paths, df_portal, df_zirve, tutar_farklari, eksik_girisler, fazla_girisler):
    """Kontrol sonucu özet mesajı oluştur"""
    portal_dosya_sayisi = len(portal_paths) if isinstance(portal_paths, list) else 1
    mesaj = f"✅ Kontrol tamamlandı!\n\n"
    mesaj += f"📁 İşlenen Portal Dosyası: {portal_dosya_sayisi}\n"
    mesaj += f"📄 Toplam Portal Fatura: {len(df_portal)}\n"
    mesaj += f"📄 Toplam Zirve Fatura: {len(df_zirve)}\n\n"
    mesaj += f"📊 Tutar Farkları: {len(tutar_farklari)}\n"
    mesaj += f"⚠️ Eksik Girişler: {len(eksik_girisler)}\n"
    mesaj += f"❌ Fazla Girişler: {len(fazla_girisler)}\n\n"
    mesaj += f"Sonuç dosyası masaüstüne kaydedildi."
    return mesaj

# ----------------- Uyumsoft Kontrol Fonksiyonları -----------------

def uyumsoft_zirve_karsilastir(zirve_path, portal_paths):
    """Zirve ve Uyumsoft Portal Excel dosyalarını karşılaştırır (Zirve programı için)."""
    global SON_DOSYA_YOLU
    
    try:
        df_zirve, df_portal = _excel_dosyalari_oku(zirve_path, portal_paths)
        
        df_zirve['Fatura_No_Temiz'] = df_zirve['Fatura No'].apply(_temizle_fatura_no)
        df_portal['Fatura_No_Temiz'] = df_portal['Fatura No'].apply(_temizle_fatura_no)
        
        # Tutarları float'a çevir
        df_zirve['Tutar_TL'] = pd.to_numeric(df_zirve['Tutar TL'], errors='coerce').fillna(0)
        df_zirve['KDV_TL'] = pd.to_numeric(df_zirve['KDV TL'], errors='coerce').fillna(0)
        df_portal['Odenecek_Tutar'] = pd.to_numeric(df_portal['Ödenecek Tutar'], errors='coerce').fillna(0)
        df_portal['Toplam_KDV'] = pd.to_numeric(df_portal['Toplam KDV'], errors='coerce').fillna(0)
        
        # Set'ler oluştur
        zirve_faturalar = set(df_zirve['Fatura_No_Temiz'])
        portal_faturalar = set(df_portal['Fatura_No_Temiz'])
        
        # 1. TUTAR FARKLARI
        tutar_farklari = []
        for fatura_no in zirve_faturalar.intersection(portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            tutar_farki = abs(zirve_row['Tutar_TL'] - portal_row['Odenecek_Tutar'])
            kdv_farki = abs(zirve_row['KDV_TL'] - portal_row['Toplam_KDV'])
            
            if tutar_farki > 1 or kdv_farki > 1:
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Zirve Cari': zirve_row.get('Cari Unvanı', ''),
                    'Portal Cari': portal_row.get('Gönderici', ''),
                    'Zirve Tutar': round(zirve_row['Tutar_TL'], 2),
                    'Portal Tutar': round(portal_row['Odenecek_Tutar'], 2),
                    'Tutar Farkı': round(tutar_farki, 2),
                    'Zirve KDV': round(zirve_row['KDV_TL'], 2),
                    'Portal KDV': round(portal_row['Toplam_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Zirve Tarih': zirve_row.get('Tarih', ''),
                    'Portal Tarih': portal_row.get('Fatura Tarihi', '')
                })
        
        # 2. EKSİK GİRİŞLER
        eksik_girisler = []
        for fatura_no in (portal_faturalar - zirve_faturalar):
            if not fatura_no:
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'Gönderici VKN/TCKN': portal_row.get('Gönderici VKN/TCKN', ''),
                'Gönderici': portal_row.get('Gönderici', ''),
                'Ödenecek Tutar': round(portal_row['Odenecek_Tutar'], 2),
                'Toplam KDV': round(portal_row['Toplam_KDV'], 2),
                'Fatura Tarihi': portal_row.get('Fatura Tarihi', ''),
                'Oluşturulma Tarihi': portal_row.get('Oluşturulma Tarihi', ''),
                'Fatura Türü': portal_row.get('Fatura Tipi', ''),
                'Fatura Durumu': portal_row.get('Fatura Durumu', '')
            })
        
        # 3. FAZLA GİRİŞLER
        fazla_girisler = []
        for fatura_no in (zirve_faturalar - portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': zirve_row.get('Cari Unvanı', ''),
                'Tutar TL': round(zirve_row['Tutar_TL'], 2),
                'KDV TL': round(zirve_row['KDV_TL'], 2),
                'Tarih': zirve_row.get('Tarih', ''),
                'E.Kod': zirve_row.get('E.Kod', ''),
                'İşlem Türü': zirve_row.get('İşlem Türü', '')
            })
        
        # Excel dosyası oluştur
        sonuc_dosyasi = os.path.join(MASAUSTU, "Uyumsoft_Zirve_Kontrol_Sonuc.xlsx")
        _kontrol_sonuc_excel_yaz(sonuc_dosyasi, tutar_farklari, eksik_girisler, fazla_girisler)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        mesaj = _kontrol_ozet_mesaj_olustur(portal_paths, df_portal, df_zirve, tutar_farklari, eksik_girisler, fazla_girisler)
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        return (None, f"❌ Hata oluştu:\n{str(e)}")

# Verim Excel sütun isimleri (A'dan N'ye)
VERIM_SUTUNLAR = ['FATURA TARİH', 'FATURA NO', 'HESAP KODU', 'HESAP ADI', 'TOPLAM', 'İND', 'ARA TOPLAM', '%1', '% 8', '% 10', '% 18', '% 20', 'NET TOPLAM', 'PB']

def _verim_excel_oku(path):
    """Verim Excel dosyasını oku ve sütun isimlerini zorla ata."""
    # Dosyayı header olmadan oku
    df = pd.read_excel(path, header=None)
    
    # İlk satıra sütun isimlerini ata (A'dan N'ye kadar 14 sütun)
    if len(df.columns) >= 14:
        df.columns = VERIM_SUTUNLAR + list(df.columns[14:])
    else:
        df.columns = VERIM_SUTUNLAR[:len(df.columns)]
    
    # İlk satırı sil (çünkü orada eski sütun isimleri var)
    df = df.iloc[1:].reset_index(drop=True)
    
    return df

# Geçersiz fatura numaraları (Export özet satırları)
GECERSIZ_FATURA_NO = ['BAŞLAMA:', 'BELGE SAYISI:', 'BİTİŞ:', 'BITIŞ:', 'SÜRE:']

def _gecersiz_fatura_no_mu(fatura_no):
    """Fatura numarasının geçersiz (export özet satırı) olup olmadığını kontrol et."""
    if not fatura_no:
        return True
    fatura_str = str(fatura_no).upper()
    for gecersiz in GECERSIZ_FATURA_NO:
        if gecersiz.upper() in fatura_str:
            return True
    return False

def _verim_normalize(df):
    """Verim Excel dosyasını normalize et - sütun adlarını standartlaştır."""
    df_normalized = df.copy()
    
    # FATURA NO -> Fatura No
    if 'FATURA NO' in df.columns:
        df_normalized['Fatura No'] = df['FATURA NO']
    
    # NET TOPLAM -> Tutar TL
    if 'NET TOPLAM' in df.columns:
        df_normalized['Tutar TL'] = pd.to_numeric(df['NET TOPLAM'], errors='coerce').fillna(0)
    
    # KDV Toplamı (%1 + % 8 + % 10 + % 18 + % 20)
    kdv_sum = 0
    for col in ['%1', '% 8', '% 10', '% 18', '% 20']:
        if col in df.columns:
            kdv_sum = kdv_sum + pd.to_numeric(df[col], errors='coerce').fillna(0)
    df_normalized['KDV TL'] = kdv_sum
    
    # HESAP ADI -> Cari Unvanı
    if 'HESAP ADI' in df.columns:
        df_normalized['Cari Unvanı'] = df['HESAP ADI']
    
    # FATURA TARİH -> Tarih
    if 'FATURA TARİH' in df.columns:
        df_normalized['Tarih'] = df['FATURA TARİH']
    
    # HESAP KODU -> E.Kod
    if 'HESAP KODU' in df.columns:
        df_normalized['E.Kod'] = df['HESAP KODU']
    
    return df_normalized

def uyumsoft_verim_karsilastir(verim_paths, portal_paths):
    """Verim ve Uyumsoft Portal Excel dosyalarını karşılaştırır (Verim programı için)."""
    global SON_DOSYA_YOLU
    
    try:
        # Verim dosyalarını oku ve normalize et
        if isinstance(verim_paths, list):
            verim_dfs = [_verim_normalize(_verim_excel_oku(path)) for path in verim_paths]
            df_verim = pd.concat(verim_dfs, ignore_index=True)
        else:
            df_verim = _verim_normalize(_verim_excel_oku(verim_paths))
        
        # Portal (Uyumsoft) dosyalarını oku
        if isinstance(portal_paths, list):
            portal_dfs = [pd.read_excel(path) for path in portal_paths]
            df_portal = pd.concat(portal_dfs, ignore_index=True)
        else:
            df_portal = pd.read_excel(portal_paths)
        
        df_verim['Fatura_No_Temiz'] = df_verim['Fatura No'].apply(_temizle_fatura_no)
        df_portal['Fatura_No_Temiz'] = df_portal['Fatura No'].apply(_temizle_fatura_no)
        
        # Tutarları float'a çevir
        df_verim['Tutar_TL'] = pd.to_numeric(df_verim['Tutar TL'], errors='coerce').fillna(0)
        df_verim['KDV_TL'] = pd.to_numeric(df_verim['KDV TL'], errors='coerce').fillna(0)
        df_portal['Odenecek_Tutar'] = pd.to_numeric(df_portal['Ödenecek Miktar'], errors='coerce').fillna(0)
        
        # Portal KDV Toplamı = Toplam Kdv 1 + Toplam Kdv 8 + Toplam Kdv 18 + Toplam KDV %10 Tutar + Toplam KDV %20 Tutar
        kdv_cols = ['Toplam Kdv 1', 'Toplam Kdv 8', 'Toplam Kdv 18', 'Toplam KDV %10 Tutar', 'Toplam KDV %20 Tutar']
        df_portal['Toplam_KDV'] = 0
        for col in kdv_cols:
            if col in df_portal.columns:
                df_portal['Toplam_KDV'] = df_portal['Toplam_KDV'] + pd.to_numeric(df_portal[col], errors='coerce').fillna(0)
        
        # Set'ler oluştur
        verim_faturalar = set(df_verim['Fatura_No_Temiz'])
        portal_faturalar = set(df_portal['Fatura_No_Temiz'])
        
        # 1. TUTAR FARKLARI
        tutar_farklari = []
        for fatura_no in verim_faturalar.intersection(portal_faturalar):
            if not fatura_no:
                continue
            verim_row = df_verim[df_verim['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            tutar_farki = abs(verim_row['Tutar_TL'] - portal_row['Odenecek_Tutar'])
            kdv_farki = abs(verim_row['KDV_TL'] - portal_row['Toplam_KDV'])
            
            if tutar_farki > 1 or kdv_farki > 1:
                # Portal tarih ve saat formatını ayır
                portal_datetime = portal_row.get('Fatura Oluşturma Tarihi', '')
                portal_tarih = ''
                portal_saat = ''
                if pd.notna(portal_datetime) and portal_datetime:
                    try:
                        dt = pd.to_datetime(portal_datetime)
                        portal_tarih = dt.strftime('%d.%m.%Y')
                        portal_saat = dt.strftime('%H:%M:%S')
                    except:
                        portal_tarih = str(portal_datetime)
                
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Verim Cari': verim_row.get('Cari Unvanı', ''),
                    'Portal Cari': portal_row.get('Cari Adı', ''),
                    'Verim Tutar': round(verim_row['Tutar_TL'], 2),
                    'Portal Tutar': round(portal_row['Odenecek_Tutar'], 2),
                    'Tutar Farkı': round(tutar_farki, 2),
                    'Verim KDV': round(verim_row['KDV_TL'], 2),
                    'Portal KDV': round(portal_row['Toplam_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Verim Tarih': verim_row.get('Tarih', ''),
                    'Portal Tarih': portal_tarih,
                    'Portal Saat': portal_saat
                })
        
        # 2. EKSİK GİRİŞLER
        eksik_girisler = []
        for fatura_no in (portal_faturalar - verim_faturalar):
            if not fatura_no:
                continue
            # Geçersiz fatura numaralarını atla (export özet satırları)
            if _gecersiz_fatura_no_mu(fatura_no):
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'Cari VKN/TCKN': portal_row.get('Cari Vkn/Tckn', ''),
                'Cari Adı': portal_row.get('Cari Adı', ''),
                'Ödenecek Tutar': round(portal_row['Odenecek_Tutar'], 2),
                'Toplam KDV': round(portal_row['Toplam_KDV'], 2),
                'Fatura Tarihi': portal_row.get('Fatura Tarihi', ''),
                'Oluşturulma Tarihi': portal_row.get('Fatura Oluşturma Tarihi', ''),
                'Fatura Türü': portal_row.get('Tür', ''),
                'Fatura Durumu': portal_row.get('Durum', '')
            })
        
        # 3. FAZLA GİRİŞLER
        fazla_girisler = []
        for fatura_no in (verim_faturalar - portal_faturalar):
            if not fatura_no:
                continue
            verim_row = df_verim[df_verim['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': verim_row.get('Cari Unvanı', ''),
                'Tutar TL': round(verim_row['Tutar_TL'], 2),
                'KDV TL': round(verim_row['KDV_TL'], 2),
                'Tarih': verim_row.get('Tarih', ''),
                'E.Kod': verim_row.get('E.Kod', '')
            })
        
        # Excel dosyası oluştur
        sonuc_dosyasi = os.path.join(MASAUSTU, "Uyumsoft_Verim_Kontrol_Sonuc.xlsx")
        _kontrol_sonuc_excel_yaz(sonuc_dosyasi, tutar_farklari, eksik_girisler, fazla_girisler)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        mesaj = _kontrol_ozet_mesaj_olustur(portal_paths, df_portal, df_verim, tutar_farklari, eksik_girisler, fazla_girisler)
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        return (None, f"❌ Hata oluştu:\n{str(e)}")

# ----------------- İzibiz Kontrol Fonksiyonu -----------------

def izibiz_karsilastir(zirve_path, portal_paths):
    """Zirve ve İzibiz Portal Excel dosyalarını karşılaştırır.
    
    Alış ve satış faturalarını otomatik tespit eder.
    """
    global SON_DOSYA_YOLU
    
    try:
        # Zirve dosyasını oku
        df_zirve = pd.read_excel(zirve_path)
        df_zirve['Fatura_No_Temiz'] = df_zirve['Fatura No'].apply(_temizle_fatura_no)
        df_zirve['Tutar_TL'] = pd.to_numeric(df_zirve['Tutar TL'], errors='coerce').fillna(0)
        df_zirve['KDV_TL'] = pd.to_numeric(df_zirve['KDV TL'], errors='coerce').fillna(0)
        
        # Portal dosyalarını oku ve normalize et
        portal_dfs = []
        if isinstance(portal_paths, list):
            for path in portal_paths:
                df = pd.read_excel(path)
                df = _izibiz_normalize_portal(df)
                portal_dfs.append(df)
            df_portal = pd.concat(portal_dfs, ignore_index=True)
        else:
            df_portal = pd.read_excel(portal_paths)
            df_portal = _izibiz_normalize_portal(df_portal)
        
        # Fatura_No_Temiz normalize fonksiyonunda oluşturuldu
        
        # Set'ler oluştur
        zirve_faturalar = set(df_zirve['Fatura_No_Temiz']) - {''}
        portal_faturalar = set(df_portal['Fatura_No_Temiz']) - {''}
        
        # 1. TUTAR FARKLARI
        tutar_farklari = []
        for fatura_no in zirve_faturalar.intersection(portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            tutar_farki = abs(zirve_row['Tutar_TL'] - portal_row['Portal_Tutar'])
            kdv_farki = abs(zirve_row['KDV_TL'] - portal_row['Portal_KDV'])
            
            if tutar_farki > 1 or kdv_farki > 1:
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Zirve Cari': zirve_row.get('Cari Unvanı', ''),
                    'Portal Cari': portal_row.get('Portal_Cari', ''),
                    'Zirve Tutar': round(zirve_row['Tutar_TL'], 2),
                    'Portal Tutar': round(portal_row['Portal_Tutar'], 2),
                    'Tutar Farkı': round(tutar_farki, 2),
                    'Zirve KDV': round(zirve_row['KDV_TL'], 2),
                    'Portal KDV': round(portal_row['Portal_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Zirve Tarih': zirve_row.get('Tarih', ''),
                    'Portal Tarih': portal_row.get('Portal_Tarih', '')
                })
        
        # 2. EKSİK GİRİŞLER
        eksik_girisler = []
        for fatura_no in (portal_faturalar - zirve_faturalar):
            if not fatura_no:
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'VKN/TCKN': portal_row.get('Portal_VKN', ''),
                'Cari': portal_row.get('Portal_Cari', ''),
                'Ödenecek Tutar': round(portal_row['Portal_Tutar'], 2),
                'Toplam KDV': round(portal_row['Portal_KDV'], 2),
                'Fatura Tarihi': portal_row.get('Portal_Tarih', ''),
                'Fatura Durumu': portal_row.get('Portal_Durum', '')
            })
        
        # 3. FAZLA GİRİŞLER
        fazla_girisler = []
        for fatura_no in (zirve_faturalar - portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': zirve_row.get('Cari Unvanı', ''),
                'Tutar TL': round(zirve_row['Tutar_TL'], 2),
                'KDV TL': round(zirve_row['KDV_TL'], 2),
                'Tarih': zirve_row.get('Tarih', ''),
                'E.Kod': zirve_row.get('E.Kod', ''),
                'İşlem Türü': zirve_row.get('İşlem Türü', '')
            })
        
        # Excel dosyası oluştur
        sonuc_dosyasi = os.path.join(MASAUSTU, "Izibiz_Kontrol_Sonuc.xlsx")
        _kontrol_sonuc_excel_yaz(sonuc_dosyasi, tutar_farklari, eksik_girisler, fazla_girisler)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        mesaj = _kontrol_ozet_mesaj_olustur(portal_paths, df_portal, df_zirve, tutar_farklari, eksik_girisler, fazla_girisler)
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        import traceback
        return (None, f"❌ Hata oluştu:\n{str(e)}\n\n{traceback.format_exc()}")

# ----------------- Vega Kontrol Fonksiyonu -----------------

def vega_karsilastir(zirve_path, portal_paths):
    """Zirve ve Vega Portal Excel dosyalarını karşılaştırır."""
    global SON_DOSYA_YOLU
    
    try:
        df_zirve, df_portal = _excel_dosyalari_oku(zirve_path, portal_paths)
        
        # Tutarları sayıya çevir (Vega'ya özgü format)
        def parse_tutar(value):
            if pd.isna(value):
                return 0
            if isinstance(value, (int, float)):
                return float(value)
            return float(str(value).replace('.', '').replace(',', '.'))
        
        df_zirve['Fatura_No_Temiz'] = df_zirve['Fatura No'].apply(_temizle_fatura_no)
        df_portal['Fatura_No_Temiz'] = df_portal['Fatura No'].apply(_temizle_fatura_no)
        
        df_zirve['Toplam_TL'] = df_zirve['Tutar TL'].apply(parse_tutar)
        df_zirve['KDV_TL'] = df_zirve['KDV TL'].apply(parse_tutar)
        
        # Portal tutarları - farklı sütun isimleri olabilir
        def get_portal_toplam(row):
            if 'Ödenecek' in df_portal.columns and pd.notna(row.get('Ödenecek')):
                return parse_tutar(row['Ödenecek'])
            elif 'Toplam' in df_portal.columns and pd.notna(row.get('Toplam')):
                return parse_tutar(row['Toplam'])
            return 0
        
        df_portal['Portal_KDV'] = df_portal['Vergi'].apply(parse_tutar)
        df_portal['Portal_Toplam'] = df_portal.apply(get_portal_toplam, axis=1)
        
        # Set'ler oluştur
        zirve_faturalar = set(df_zirve['Fatura_No_Temiz']) - {''}
        portal_faturalar = set(df_portal['Fatura_No_Temiz']) - {''}
        
        # 1. TUTAR FARKLARI
        tutar_farklari = []
        for fatura_no in zirve_faturalar.intersection(portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            toplam_farki = abs(zirve_row['Toplam_TL'] - portal_row['Portal_Toplam'])
            kdv_farki = abs(zirve_row['KDV_TL'] - portal_row['Portal_KDV'])
            
            if toplam_farki > 1 or kdv_farki > 1:
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Zirve Cari': zirve_row.get('Cari Unvanı', ''),
                    'Portal Cari': portal_row.get('Unvan', ''),
                    'Zirve Tutar': round(zirve_row['Toplam_TL'], 2),
                    'Portal Tutar': round(portal_row['Portal_Toplam'], 2),
                    'Tutar Farkı': round(toplam_farki, 2),
                    'Zirve KDV': round(zirve_row['KDV_TL'], 2),
                    'Portal KDV': round(portal_row['Portal_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Zirve Tarih': zirve_row.get('Tarih', ''),
                    'Portal Tarih': portal_row.get('Tarih', '')
                })
        
        # 2. EKSİK GİRİŞLER
        # Vega tarih formatını parse et (örn: '2025.12.0101.12.2025' -> '01.12.2025')
        def parse_vega_tarih(tarih_str):
            if pd.isna(tarih_str) or tarih_str == '':
                return ''
            tarih_str = str(tarih_str)
            if len(tarih_str) >= 10 and '.' in tarih_str:
                parts = tarih_str.split('.')
                if len(parts) >= 4:
                    try:
                        return f"{parts[-3][-2:]}.{parts[-2]}.{parts[-1]}"
                    except:
                        pass
            return tarih_str
        
        eksik_girisler = []
        for fatura_no in (portal_faturalar - zirve_faturalar):
            if not fatura_no:
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            # Fatura Türü: Türü veya Fatura Türü sütunlarından al
            fatura_turu = ''
            if 'Türü' in portal_row.index and pd.notna(portal_row['Türü']):
                fatura_turu = str(portal_row['Türü'])
            elif 'Fatura Türü' in portal_row.index and pd.notna(portal_row['Fatura Türü']):
                fatura_turu = str(portal_row['Fatura Türü'])
            
            # Fatura Durumu: Kabul/Red, Red/Kabul veya Rapor Durumu sütunlarından al
            fatura_durumu = ''
            if 'Kabul/Red' in portal_row.index and pd.notna(portal_row['Kabul/Red']):
                fatura_durumu = str(portal_row['Kabul/Red'])
            elif 'Red/Kabul' in portal_row.index and pd.notna(portal_row['Red/Kabul']):
                fatura_durumu = str(portal_row['Red/Kabul'])
            elif 'Rapor Durumu' in portal_row.index and pd.notna(portal_row['Rapor Durumu']):
                fatura_durumu = str(portal_row['Rapor Durumu'])
            
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'VKN/TCKN': portal_row.get('VKN/TCKN', ''),
                'Unvan': portal_row.get('Unvan', ''),
                'Ödenecek Tutar': round(portal_row['Portal_Toplam'], 2),
                'Toplam KDV': round(portal_row['Portal_KDV'], 2),
                'Fatura Tarihi': parse_vega_tarih(portal_row.get('Tarih', '')),
                'Oluşturulma Tarihi': portal_row.get('Oluşturulma Tarihi', portal_row.get('Kayıt Tarihi', '')),
                'Fatura Türü': fatura_turu,
                'Fatura Durumu': fatura_durumu
            })
        
        # 3. FAZLA GİRİŞLER
        fazla_girisler = []
        for fatura_no in (zirve_faturalar - portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': zirve_row.get('Cari Unvanı', ''),
                'Tutar TL': round(zirve_row['Toplam_TL'], 2),
                'KDV TL': round(zirve_row['KDV_TL'], 2),
                'Tarih': zirve_row.get('Tarih', ''),
                'E.Kod': zirve_row.get('E.Kod', ''),
                'İşlem Türü': zirve_row.get('İşlem Türü', '')
            })
        
        # Excel dosyası oluştur
        sonuc_dosyasi = os.path.join(MASAUSTU, "Vega_Kontrol_Sonuc.xlsx")
        _kontrol_sonuc_excel_yaz(sonuc_dosyasi, tutar_farklari, eksik_girisler, fazla_girisler)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        mesaj = _kontrol_ozet_mesaj_olustur(portal_paths, df_portal, df_zirve, tutar_farklari, eksik_girisler, fazla_girisler)
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        import traceback
        return (None, f"❌ Hata oluştu:\n{str(e)}\n\n{traceback.format_exc()}")

# ----------------- Hızlıbilişim Kontrol Fonksiyonu -----------------

def hizlibilisim_karsilastir(zirve_path, portal_paths, fatura_yonu='satis'):
    """Zirve ve Hızlıbilişim Portal Excel dosyalarını karşılaştırır.
    
    Args:
        zirve_path: Zirve Excel dosya yolu
        portal_paths: Portal Excel dosya yolu(ları)
        fatura_yonu: 'alis' veya 'satis'
            - alis: Gelen faturalar (GondericiUnvan/Vkn kullanılır)
            - satis: E-Arşiv + Giden faturalar (AliciUnvan/Vkn kullanılır)
    """
    global SON_DOSYA_YOLU
    
    try:
        # Zirve dosyasını oku
        df_zirve = pd.read_excel(zirve_path)
        df_zirve['Fatura_No_Temiz'] = df_zirve['Fatura No'].apply(_temizle_fatura_no)
        df_zirve['Tutar_TL'] = pd.to_numeric(df_zirve['Tutar TL'], errors='coerce').fillna(0)
        df_zirve['KDV_TL'] = pd.to_numeric(df_zirve['KDV TL'], errors='coerce').fillna(0)
        # Dövizli faturalar için Tutar Dvz ve Kdv Dvz sütunlarını oku
        df_zirve['Tutar_Dvz'] = pd.to_numeric(df_zirve.get('Tutar Dvz', 0), errors='coerce').fillna(0)
        df_zirve['KDV_Dvz'] = pd.to_numeric(df_zirve.get('Kdv Dvz', 0), errors='coerce').fillna(0)
        
        # Portal dosyalarını oku ve normalize et
        portal_dfs = []
        if isinstance(portal_paths, list):
            for path in portal_paths:
                df = pd.read_excel(path)
                df = _hizlibilisim_normalize_portal(df, fatura_yonu)
                portal_dfs.append(df)
            df_portal = pd.concat(portal_dfs, ignore_index=True)
        else:
            df_portal = pd.read_excel(portal_paths)
            df_portal = _hizlibilisim_normalize_portal(df_portal, fatura_yonu)
        
        df_portal['Fatura_No_Temiz'] = df_portal['FaturaNo'].apply(_temizle_fatura_no)
        
        # Set'ler oluştur
        zirve_faturalar = set(df_zirve['Fatura_No_Temiz']) - {''}
        portal_faturalar = set(df_portal['Fatura_No_Temiz']) - {''}
        
        # 1. TUTAR FARKLARI
        tutar_farklari = []
        for fatura_no in zirve_faturalar.intersection(portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            # ParaBirimi'ne göre Zirve'den doğru sütunları kullan
            para_birimi = str(portal_row.get('Portal_ParaBirimi', 'TRY')).upper()
            if para_birimi != 'TRY' and para_birimi != '':
                # Dövizli fatura - Tutar Dvz ve Kdv Dvz kullan
                zirve_tutar = zirve_row['Tutar_Dvz']
                zirve_kdv = zirve_row['KDV_Dvz']
            else:
                # TRY fatura - Tutar TL ve KDV TL kullan
                zirve_tutar = zirve_row['Tutar_TL']
                zirve_kdv = zirve_row['KDV_TL']
            
            tutar_farki = abs(zirve_tutar - portal_row['Portal_Tutar'])
            kdv_farki = abs(zirve_kdv - portal_row['Portal_KDV'])
            
            if tutar_farki > 1 or kdv_farki > 1:
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Zirve Cari': zirve_row.get('Cari Unvanı', ''),
                    'Portal Cari': portal_row.get('Portal_Cari', ''),
                    'Zirve Tutar': round(zirve_tutar, 2),
                    'Portal Tutar': round(portal_row['Portal_Tutar'], 2),
                    'Tutar Farkı': round(tutar_farki, 2),
                    'Zirve KDV': round(zirve_kdv, 2),
                    'Portal KDV': round(portal_row['Portal_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Zirve Tarih': zirve_row.get('Tarih', ''),
                    'Portal Tarih': portal_row.get('FaturaTarihi', ''),
                    'Para Birimi': para_birimi
                })
        
        # 2. EKSİK GİRİŞLER
        eksik_girisler = []
        for fatura_no in (portal_faturalar - zirve_faturalar):
            if not fatura_no:
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'VKN/TCKN': portal_row.get('Portal_VKN', ''),
                'Cari': portal_row.get('Portal_Cari', ''),
                'Ödenecek Tutar': round(portal_row['Portal_Tutar'], 2),
                'Toplam KDV': round(portal_row['Portal_KDV'], 2),
                'Fatura Tarihi': portal_row.get('FaturaTarihi', ''),
                'Oluşturulma Tarihi': portal_row.get('FaturaAlinmaTarihi', ''),
                'Fatura Türü': portal_row.get('FaturaTipi', ''),
                'Fatura Durumu': portal_row.get('Portal_Durum', '')
            })
        
        # 3. FAZLA GİRİŞLER
        fazla_girisler = []
        for fatura_no in (zirve_faturalar - portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': zirve_row.get('Cari Unvanı', ''),
                'Tutar TL': round(zirve_row['Tutar_TL'], 2),
                'KDV TL': round(zirve_row['KDV_TL'], 2),
                'Tarih': zirve_row.get('Tarih', ''),
                'E.Kod': zirve_row.get('E.Kod', ''),
                'İşlem Türü': zirve_row.get('İşlem Türü', '')
            })
        
        # Excel dosyası oluştur
        dosya_adi = "Hizlibilisim_Alis_Kontrol_Sonuc.xlsx" if fatura_yonu == 'alis' else "Hizlibilisim_Satis_Kontrol_Sonuc.xlsx"
        sonuc_dosyasi = os.path.join(MASAUSTU, dosya_adi)
        _kontrol_sonuc_excel_yaz(sonuc_dosyasi, tutar_farklari, eksik_girisler, fazla_girisler)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        mesaj = _kontrol_ozet_mesaj_olustur(portal_paths, df_portal, df_zirve, tutar_farklari, eksik_girisler, fazla_girisler)
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        import traceback
        return (None, f"❌ Hata oluştu:\n{str(e)}\n\n{traceback.format_exc()}")

def _hizlibilisim_normalize_portal(df, fatura_yonu='satis'):
    """Hızlıbilişim portal dosyasını normalize et.
    
    Args:
        df: Portal DataFrame
        fatura_yonu: 'alis' veya 'satis'
            - alis: Gelen faturalar -> GondericiUnvan/Vkn kullanılır
            - satis: E-Arşiv + Giden -> AliciUnvan/Vkn kullanılır
    """
    # Tutar belirleme - hangi sütun varsa onu kullan
    if 'OdenecekTutar' in df.columns:
        df['Portal_Tutar'] = pd.to_numeric(df['OdenecekTutar'], errors='coerce').fillna(0)
    elif 'PayableAmount' in df.columns:
        df['Portal_Tutar'] = pd.to_numeric(df['PayableAmount'], errors='coerce').fillna(0)
    else:
        df['Portal_Tutar'] = 0
    
    # KDV hesaplama - her durumda aynı (Kdv20 + Kdv10 + Kdv1)
    kdv20 = pd.to_numeric(df['Kdv20Tutar'], errors='coerce').fillna(0) if 'Kdv20Tutar' in df.columns else pd.Series([0] * len(df))
    kdv10 = pd.to_numeric(df['Kdv10Tutar'], errors='coerce').fillna(0) if 'Kdv10Tutar' in df.columns else pd.Series([0] * len(df))
    kdv1 = pd.to_numeric(df['Kdv1Tutar'], errors='coerce').fillna(0) if 'Kdv1Tutar' in df.columns else pd.Series([0] * len(df))
    df['Portal_KDV'] = kdv20 + kdv10 + kdv1
    
    # Para birimi
    df['Portal_ParaBirimi'] = df['ParaBirimi'] if 'ParaBirimi' in df.columns else 'TRY'
    
    # VKN ve Cari belirleme - fatura yönüne göre
    if fatura_yonu == 'alis':
        # Alış: Gelen faturalar -> GondericiUnvan/Vkn
        df['Portal_VKN'] = df['GondericiVkn'] if 'GondericiVkn' in df.columns else ''
        df['Portal_Cari'] = df['GondericiUnvan'] if 'GondericiUnvan' in df.columns else ''
    else:
        # Satış: Giden faturalar veya E-Arşiv -> AliciUnvan/Vkn
        df['Portal_VKN'] = df['AliciVkn'] if 'AliciVkn' in df.columns else ''
        df['Portal_Cari'] = df['AliciUnvan'] if 'AliciUnvan' in df.columns else ''
    
    # Durum sütunu - hem alış hem satış için
    if 'IptalDurumu' in df.columns:
        df['Portal_Durum'] = df['IptalDurumu']
    elif 'Durum' in df.columns:
        df['Portal_Durum'] = df['Durum']
    else:
        df['Portal_Durum'] = ''
    
    return df

def _logo_normalize_portal(df):
    """Logo portal dosyasını normalize et.
    
    Satış ve alış faturalarını otomatik tespit eder ve standart formata dönüştürür.
    """
    # Otomatik satış/alış tespiti
    # Gönderici VKN varsa alış, Alıcı Vkn/Tckn varsa satış
    is_alis = 'Gönderici VKN' in df.columns or 'Gönderici Adı' in df.columns
    
    # Fatura No
    df['FaturaNo'] = df['Fatura No'] if 'Fatura No' in df.columns else ''
    
    # Tutar ve KDV
    df['Portal_Tutar'] = pd.to_numeric(df['Toplam Tutar'], errors='coerce').fillna(0) if 'Toplam Tutar' in df.columns else 0
    df['Portal_KDV'] = pd.to_numeric(df['KDV Toplamı'], errors='coerce').fillna(0) if 'KDV Toplamı' in df.columns else 0
    
    # Para Birimi
    df['Portal_ParaBirimi'] = df['Para Birimi'] if 'Para Birimi' in df.columns else 'TRY'
    
    # VKN ve Cari (alış/satış'a göre)
    if is_alis:
        # Alış faturası - Gönderici bilgileri
        df['Portal_VKN'] = df['Gönderici VKN'] if 'Gönderici VKN' in df.columns else ''
        df['Portal_Cari'] = df['Gönderici Adı'] if 'Gönderici Adı' in df.columns else ''
    else:
        # Satış faturası - Alıcı bilgileri
        df['Portal_VKN'] = df['Alıcı Vkn/Tckn'] if 'Alıcı Vkn/Tckn' in df.columns else ''
        df['Portal_Cari'] = df['Alıcı Ünvanı'] if 'Alıcı Ünvanı' in df.columns else ''
    
    # Tarih
    df['FaturaTarihi'] = df['Fatura Tarihi'] if 'Fatura Tarihi' in df.columns else ''
    
    # Durum - Öncelik sırası: İptal/İtiraz Durumu > Durumu > Durum
    df['Portal_Durum'] = ''
    if 'İptal/İtiraz Durumu' in df.columns:
        # İptal/İtiraz Durumu varsa önce ona bak
        df['Portal_Durum'] = df['İptal/İtiraz Durumu'].fillna('')
        # Boş olanlar için diğer sütunlara bak
        if 'Durumu' in df.columns:
            mask = (df['Portal_Durum'] == '') | (df['Portal_Durum'].isna())
            df.loc[mask, 'Portal_Durum'] = df.loc[mask, 'Durumu'].fillna('')
        elif 'Durum' in df.columns:
            mask = (df['Portal_Durum'] == '') | (df['Portal_Durum'].isna())
            df.loc[mask, 'Portal_Durum'] = df.loc[mask, 'Durum'].fillna('')
    elif 'Durumu' in df.columns:
        df['Portal_Durum'] = df['Durumu'].fillna('')
    elif 'Durum' in df.columns:
        df['Portal_Durum'] = df['Durum'].fillna('')
    
    return df

def _izibiz_normalize_portal(df):
    """İzibiz portal dosyasını normalize et.
    
    Alış ve satış faturalarını otomatik tespit eder:
    - Alış: Fatura No + VKN/TCKN + Unvan → Gönderici bilgisi
    - Satış: Belge No + Gönderici Unvan + Alıcı Unvan → Alıcı bilgisi
    """
    # Otomatik alış/satış tespiti
    is_alis = ('VKN/TCKN' in df.columns and 'Unvan' in df.columns 
               and 'Fatura No' in df.columns and 'Belge No' not in df.columns)
    
    if is_alis:
        # ALIŞ FATURASI (GELEN)
        df['Fatura_No_Temiz'] = df['Fatura No']
        df['Portal_Tutar'] = pd.to_numeric(df['Tutar'], errors='coerce').fillna(0)
        df['Portal_KDV'] = 0  # Alış faturalarında ayrı KDV sütunu yok, tutara dahil
        df['Portal_VKN'] = df['VKN/TCKN']
        df['Portal_Cari'] = df['Unvan']
        df['Portal_Tarih'] = df['Tarih'] if 'Tarih' in df.columns else ''
        df['Portal_Durum'] = df['Durum'] if 'Durum' in df.columns else ''
    else:
        # SATIŞ FATURASI (GİDEN)
        df['Fatura_No_Temiz'] = df['Belge No']
        df['Portal_Tutar'] = pd.to_numeric(df['Ödenecek Tutar'], errors='coerce').fillna(0)
        df['Portal_KDV'] = pd.to_numeric(df['Toplam KDV'], errors='coerce').fillna(0)
        # Satış faturasında ALICI bilgisi kullanılmalı (müşteri)
        df['Portal_VKN'] = df.get('Alıcı VKN/TCKN', '')
        df['Portal_Cari'] = df.get('Alıcı Unvan', '')
        df['Portal_Tarih'] = df['Tarih'] if 'Tarih' in df.columns else ''
        df['Portal_Durum'] = df['Durum'] if 'Durum' in df.columns else ''
    
    return df

# ----------------- Logo + Hızlıbilişim Kombine Kontrol Fonksiyonu -----------------

def logo_hizli_kombine_karsilastir(zirve_path, logo_portal_paths, hizli_portal_paths):
    """Zirve ile Logo ve Hızlıbilişim Portal dosyalarını kombine karşılaştırır.
    
    Args:
        zirve_path: Zirve Excel dosya yolu
        logo_portal_paths: Logo portal dosya yolları listesi (boş liste olabilir)
        hizli_portal_paths: Hızlıbilişim portal dosya yolları listesi (boş liste olabilir)
    """
    global SON_DOSYA_YOLU
    
    try:
        # Zirve dosyasını oku
        df_zirve = pd.read_excel(zirve_path)
        df_zirve['Fatura_No_Temiz'] = df_zirve['Fatura No'].apply(_temizle_fatura_no)
        df_zirve['Tutar_TL'] = pd.to_numeric(df_zirve['Tutar TL'], errors='coerce').fillna(0)
        df_zirve['KDV_TL'] = pd.to_numeric(df_zirve['KDV TL'], errors='coerce').fillna(0)
        # Dövizli faturalar için
        df_zirve['Tutar_Dvz'] = pd.to_numeric(df_zirve.get('Tutar Dvz', 0), errors='coerce').fillna(0)
        df_zirve['KDV_Dvz'] = pd.to_numeric(df_zirve.get('Kdv Dvz', 0), errors='coerce').fillna(0)
        
        # Logo portal dosyalarını normalize et ve birleştir
        all_portal_dfs = []
        
        if logo_portal_paths:
            for path in logo_portal_paths:
                df = pd.read_excel(path)
                df = _logo_normalize_portal(df)
                all_portal_dfs.append(df)
        
        # Hızlıbilişim portal dosyalarını normalize et ve birleştir
        if hizli_portal_paths:
            for path in hizli_portal_paths:
                df = pd.read_excel(path)
                # Otomatik alış/satış tespiti: PayableAmount varsa Gelen/Giden, OdenecekTutar varsa E-Arşiv
                if 'PayableAmount' in df.columns:
                    # Gelen/Giden format
                    fatura_yonu = 'alis' if 'GondericiVkn' in df.columns else 'satis'
                else:
                    # E-Arşiv format (her zaman satış)
                    fatura_yonu = 'satis'
                df = _hizlibilisim_normalize_portal(df, fatura_yonu)
                all_portal_dfs.append(df)
        
        # Tüm portal dosyalarını birleştir
        if not all_portal_dfs:
            return (None, "❌ Hata: En az bir Logo veya Hızlıbilişim portal dosyası seçmelisiniz.")
        
        df_portal = pd.concat(all_portal_dfs, ignore_index=True)
        df_portal['Fatura_No_Temiz'] = df_portal['FaturaNo'].apply(_temizle_fatura_no)
        
        # Set'ler oluştur
        zirve_faturalar = set(df_zirve['Fatura_No_Temiz']) - {''}
        portal_faturalar = set(df_portal['Fatura_No_Temiz']) - {''}
        
        # 1. TUTAR FARKLARI
        tutar_farklari = []
        for fatura_no in zirve_faturalar.intersection(portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            # Para birimine göre Zirve'den doğru sütunları kullan
            para_birimi = str(portal_row.get('Portal_ParaBirimi', 'TRY')).upper()
            if para_birimi != 'TRY' and para_birimi != '':
                # Dövizli fatura
                zirve_tutar = zirve_row['Tutar_Dvz']
                zirve_kdv = zirve_row['KDV_Dvz']
            else:
                # TRY fatura
                zirve_tutar = zirve_row['Tutar_TL']
                zirve_kdv = zirve_row['KDV_TL']
            
            tutar_farki = abs(zirve_tutar - portal_row['Portal_Tutar'])
            kdv_farki = abs(zirve_kdv - portal_row['Portal_KDV'])
            
            if tutar_farki > 1 or kdv_farki > 1:
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Zirve Cari': zirve_row.get('Cari Unvanı', ''),
                    'Portal Cari': portal_row.get('Portal_Cari', ''),
                    'Zirve Tutar': round(zirve_tutar, 2),
                    'Portal Tutar': round(portal_row['Portal_Tutar'], 2),
                    'Tutar Farkı': round(tutar_farki, 2),
                    'Zirve KDV': round(zirve_kdv, 2),
                    'Portal KDV': round(portal_row['Portal_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Zirve Tarih': zirve_row.get('Tarih', ''),
                    'Portal Tarih': portal_row.get('FaturaTarihi', ''),
                    'Para Birimi': para_birimi
                })
        
        # 2. EKSİK GİRİŞLER (Portalda var, Zirve'de yok)
        eksik_girisler = []
        for fatura_no in (portal_faturalar - zirve_faturalar):
            if not fatura_no:
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'VKN/TCKN': portal_row.get('Portal_VKN', ''),
                'Cari': portal_row.get('Portal_Cari', ''),
                'Ödenecek Tutar': round(portal_row['Portal_Tutar'], 2),
                'Toplam KDV': round(portal_row['Portal_KDV'], 2),
                'Fatura Tarihi': portal_row.get('FaturaTarihi', ''),
                'Para Birimi': portal_row.get('Portal_ParaBirimi', 'TRY'),
                'Fatura Durumu': portal_row.get('Portal_Durum', '')
            })
        
        # 3. FAZLA GİRİŞLER (Zirve'de var, portalda yok)
        fazla_girisler = []
        for fatura_no in (zirve_faturalar - portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': zirve_row.get('Cari Unvanı', ''),
                'Tutar TL': round(zirve_row['Tutar_TL'], 2),
                'KDV TL': round(zirve_row['KDV_TL'], 2),
                'Tarih': zirve_row.get('Tarih', ''),
                'E.Kod': zirve_row.get('E.Kod', ''),
                'İşlem Türü': zirve_row.get('İşlem Türü', '')
            })
        
        # Excel dosyası oluştur
        sonuc_dosyasi = os.path.join(MASAUSTU, "Logo_Zirve_Kontrol_Sonuc.xlsx")
        _kontrol_sonuc_excel_yaz(sonuc_dosyasi, tutar_farklari, eksik_girisler, fazla_girisler)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        
        # Standart mesaj oluştur (diğer kontroller gibi)
        mesaj = _kontrol_ozet_mesaj_olustur(
            logo_portal_paths + hizli_portal_paths,  # Tüm portal dosyaları
            df_portal, 
            df_zirve, 
            tutar_farklari, 
            eksik_girisler, 
            fazla_girisler
        )
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        import traceback
        return (None, f"❌ Hata oluştu:\n{str(e)}\n\n{traceback.format_exc()}")

# ----------------- Portal Tipi Tespit Fonksiyonu -----------------

def _tespit_portal_tipi(df):
    """Portal Excel dosyasının tipini tespit eder (Hepsiburada veya Trendyol).
    
    Args:
        df: Portal DataFrame
        
    Returns:
        'trendyol' veya 'hepsiburada'
    """
    sutunlar = df.columns.tolist()
    
    # Trendyol tespiti: "Gönderici Ünvan" VE "KDV Toplamı" varsa
    if 'Gönderici Ünvan' in sutunlar and 'KDV Toplamı' in sutunlar:
        return 'trendyol'
    # Hepsiburada tespiti: ("Gönderici" VEYA "Alıcı") VE "Toplam KDV" varsa
    elif ('Gönderici' in sutunlar or 'Alıcı' in sutunlar) and 'Toplam KDV' in sutunlar:
        return 'hepsiburada'
    else:
        return 'bilinmiyor'

# ----------------- Hepsiburada/Trendyol Kontrol Fonksiyonu -----------------

def hb_ve_trn_karsilastir(zirve_path, portal_paths):
    """Zirve ve Hepsiburada/Trendyol Portal Excel dosyalarını karşılaştırır.
    Portal tipi ve Gelen/Giden otomatik tespit edilir.
    
    Args:
        zirve_path: Zirve Excel dosya yolu
        portal_paths: Portal Excel dosya yolu(ları)
    """
    global SON_DOSYA_YOLU
    
    try:
        # Zirve dosyasını oku
        df_zirve = pd.read_excel(zirve_path)
        df_zirve['Fatura_No_Temiz'] = df_zirve['Fatura No'].apply(_temizle_fatura_no)
        df_zirve['Tutar_TL'] = pd.to_numeric(df_zirve['Tutar TL'], errors='coerce').fillna(0)
        df_zirve['KDV_TL'] = pd.to_numeric(df_zirve['KDV TL'], errors='coerce').fillna(0)
        
        # Portal dosyalarını oku ve normalize et (otomatik tespit)
        portal_dfs = []
        if isinstance(portal_paths, list):
            for path in portal_paths:
                df = pd.read_excel(path)
                portal_tipi = _tespit_portal_tipi(df)
                if portal_tipi == 'trendyol':
                    df = _trendyol_normalize_portal(df)
                elif portal_tipi == 'hepsiburada':
                    df = _hepsiburada_normalize_portal(df)
                else:
                    raise ValueError(f"Bilinmeyen portal tipi: {portal_tipi}")
                portal_dfs.append(df)
            df_portal = pd.concat(portal_dfs, ignore_index=True)
        else:
            df_portal = pd.read_excel(portal_paths)
            portal_tipi = _tespit_portal_tipi(df_portal)
            if portal_tipi == 'trendyol':
                df_portal = _trendyol_normalize_portal(df_portal)
            elif portal_tipi == 'hepsiburada':
                df_portal = _hepsiburada_normalize_portal(df_portal)
            else:
                raise ValueError(f"Bilinmeyen portal tipi: {portal_tipi}")
        
        df_portal['Fatura_No_Temiz'] = df_portal['FaturaNo'].apply(_temizle_fatura_no)
        
        # Set'ler oluştur
        zirve_faturalar = set(df_zirve['Fatura_No_Temiz']) - {''}
        portal_faturalar = set(df_portal['Fatura_No_Temiz']) - {''}
        
        # 1. TUTAR FARKLARI
        tutar_farklari = []
        for fatura_no in zirve_faturalar.intersection(portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            tutar_farki = abs(zirve_row['Tutar_TL'] - portal_row['Portal_Tutar'])
            kdv_farki = abs(zirve_row['KDV_TL'] - portal_row['Portal_KDV'])
            
            if tutar_farki > 1 or kdv_farki > 1:
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Portal Kaynak': portal_row.get('Portal_Kaynak', ''),
                    'Zirve Cari': zirve_row.get('Cari Unvanı', ''),
                    'Portal Gönderici': portal_row.get('Portal_Gonderici', ''),
                    'Portal Alıcı': portal_row.get('Portal_Alici', ''),
                    'Zirve Tutar': round(zirve_row['Tutar_TL'], 2),
                    'Portal Tutar': round(portal_row['Portal_Tutar'], 2),
                    'Tutar Farkı': round(tutar_farki, 2),
                    'Zirve KDV': round(zirve_row['KDV_TL'], 2),
                    'Portal KDV': round(portal_row['Portal_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Zirve Tarih': zirve_row.get('Tarih', ''),
                    'Portal Tarih': portal_row.get('FaturaTarihi', '')
                })
        
        # 2. EKSİK GİRİŞLER
        eksik_girisler = []
        for fatura_no in (portal_faturalar - zirve_faturalar):
            if not fatura_no:
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'Portal Kaynak': portal_row.get('Portal_Kaynak', ''),
                'Gönderici': portal_row.get('Portal_Gonderici', ''),
                'Alıcı': portal_row.get('Portal_Alici', ''),
                'VKN/TCKN': portal_row.get('Portal_VKN', ''),
                'Ödenecek Tutar': round(portal_row['Portal_Tutar'], 2),
                'Toplam KDV': round(portal_row['Portal_KDV'], 2),
                'Fatura Tarihi': portal_row.get('FaturaTarihi', ''),
                'Oluşturulma Tarihi': portal_row.get('OlusturmaTarihi', ''),
                'Fatura Türü': portal_row.get('FaturaTipi', ''),
                'Fatura Durumu': portal_row.get('Portal_Durum', '')
            })
        
        # 3. FAZLA GİRİŞLER
        fazla_girisler = []
        for fatura_no in (zirve_faturalar - portal_faturalar):
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': zirve_row.get('Cari Unvanı', ''),
                'Tutar TL': round(zirve_row['Tutar_TL'], 2),
                'KDV TL': round(zirve_row['KDV_TL'], 2),
                'Tarih': zirve_row.get('Tarih', ''),
                'E.Kod': zirve_row.get('E.Kod', ''),
                'İşlem Türü': zirve_row.get('İşlem Türü', '')
            })
        
        # Excel dosyası oluştur
        sonuc_dosyasi = os.path.join(MASAUSTU, "HB_TRN_Kontrol_Sonuc.xlsx")
        _kontrol_sonuc_excel_yaz(sonuc_dosyasi, tutar_farklari, eksik_girisler, fazla_girisler)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        mesaj = _kontrol_ozet_mesaj_olustur(portal_paths, df_portal, df_zirve, tutar_farklari, eksik_girisler, fazla_girisler)
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        import traceback
        return (None, f"❌ Hata oluştu:\n{str(e)}\n\n{traceback.format_exc()}")

def _hepsiburada_normalize_portal(df):
    """Hepsiburada portal dosyasını normalize et.
    Gelen (Gönderici) veya Giden (Alıcı) otomatik tespit edilir.
    
    Args:
        df: Portal DataFrame
    """
    # Portal kaynağı
    df['Portal_Kaynak'] = 'Hepsiburada'
    
    # Fatura No
    df['FaturaNo'] = df['Fatura No'] if 'Fatura No' in df.columns else ''
    
    # Tutarlar
    df['Portal_Tutar'] = pd.to_numeric(df['Ödenecek Tutar'], errors='coerce').fillna(0)
    df['Portal_KDV'] = pd.to_numeric(df['Toplam KDV'], errors='coerce').fillna(0)
    
    # Tarihler
    df['FaturaTarihi'] = df['Fatura Tarihi'] if 'Fatura Tarihi' in df.columns else ''
    df['OlusturmaTarihi'] = df['Oluşturulma Tarihi'] if 'Oluşturulma Tarihi' in df.columns else ''
    
    # Fatura bilgileri
    df['FaturaTipi'] = df['Fatura Tipi'] if 'Fatura Tipi' in df.columns else ''
    df['Portal_Durum'] = df['Fatura Durumu'] if 'Fatura Durumu' in df.columns else ''
    
    # Cari bilgileri - Otomatik tespit (Gönderici varsa Gelen, Alıcı varsa Giden)
    if 'Gönderici' in df.columns:
        # Gelen faturalar -> Gönderici
        df['Portal_VKN'] = df['Gönderici VKN/TCKN'] if 'Gönderici VKN/TCKN' in df.columns else ''
        df['Portal_Gonderici'] = df['Gönderici']
        df['Portal_Alici'] = ''
        df['Portal_Cari'] = df['Gönderici']  # Geriye uyumluluk için
    elif 'Alıcı' in df.columns:
        # Giden faturalar -> Alıcı
        df['Portal_VKN'] = df['Alıcı VKN/TCKN'] if 'Alıcı VKN/TCKN' in df.columns else ''
        df['Portal_Gonderici'] = ''
        df['Portal_Alici'] = df['Alıcı']
        df['Portal_Cari'] = df['Alıcı']  # Geriye uyumluluk için
    else:
        # Ne Gönderici ne de Alıcı sütunu yok
        df['Portal_VKN'] = ''
        df['Portal_Gonderici'] = ''
        df['Portal_Alici'] = ''
        df['Portal_Cari'] = ''
    
    return df

def _trendyol_normalize_portal(df):
    """Trendyol portal dosyasını normalize et.
    İki farklı format desteklenir:
    - Format 1: Durum, Zarf UUID
    - Format 2: Statü, Rapor Durumu
    
    Args:
        df: Portal DataFrame
    """
    # Portal kaynağı
    df['Portal_Kaynak'] = 'Trendyol'
    
    # Fatura No
    df['FaturaNo'] = df['Fatura No'] if 'Fatura No' in df.columns else ''
    
    # Tutarlar
    df['Portal_Tutar'] = pd.to_numeric(df['Ödenecek Tutar'], errors='coerce').fillna(0)
    df['Portal_KDV'] = pd.to_numeric(df['KDV Toplamı'], errors='coerce').fillna(0)
    
    # Tarihler
    df['FaturaTarihi'] = df['Fatura Tarihi'] if 'Fatura Tarihi' in df.columns else ''
    df['OlusturmaTarihi'] = df['Oluşturulma Tarihi'] if 'Oluşturulma Tarihi' in df.columns else ''
    
    # Fatura bilgileri
    df['FaturaTipi'] = df['Fatura Tipi'] if 'Fatura Tipi' in df.columns else ''
    
    # Fatura durumu - iki farklı format
    if 'Durum' in df.columns:
        df['Portal_Durum'] = df['Durum']
    elif 'Statü' in df.columns:
        df['Portal_Durum'] = df['Statü']
    else:
        df['Portal_Durum'] = ''
    
    # Cari bilgileri - Hem Gönderici hem Alıcı var
    df['Portal_Gonderici'] = df['Gönderici Ünvan'] if 'Gönderici Ünvan' in df.columns else ''
    df['Portal_Alici'] = df['Alıcı Ünvan'] if 'Alıcı Ünvan' in df.columns else ''
    
    # VKN bilgisi
    if 'Gönderici VKN/TCKN' in df.columns:
        df['Portal_VKN'] = df['Gönderici VKN/TCKN']
    elif 'Alıcı VKN/TCKN' in df.columns:
        df['Portal_VKN'] = df['Alıcı VKN/TCKN']
    else:
        df['Portal_VKN'] = ''
    
    # Portal_Cari - Gönderici veya Alıcı (hangisi dolu ise)
    df['Portal_Cari'] = df['Portal_Gonderici'].where(df['Portal_Gonderici'] != '', df['Portal_Alici'])
    
    return df

# ----------------- Logo XML -----------------

def xml_logo_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    global SON_DOSYA_YOLU
    veriler = []
    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            ns = XML_NS
            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)
            carikodu = root.findtext(".//cac:AccountingCustomerParty//cbc:ID", namespaces=ns)
            party = root.find(".//cac:AccountingCustomerParty//cac:Party", namespaces=ns)
            cariadi = cari_adi_al(party, ns) if party is not None else "Bilinmeyen Cari"
            doviz = root.findtext(".//cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"

            for line in root.findall(".//cac:InvoiceLine", namespaces=ns):
                miktar = int(float(line.findtext("cbc:InvoicedQuantity", default="0", namespaces=ns)))
                birim = "ADET"
                fiyat = float(line.findtext("cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext("cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext("cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns)

                iskonto_element = line.find("cac:AllowanceCharge/cbc:Amount", namespaces=ns)
                iskonto = float(iskonto_element.text.replace(",", ".")) if iskonto_element is not None else 0.0

                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except (ValueError, TypeError):
                    kdv_orani = 0

                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Satış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": birim,
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto,
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")

    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, f"{fatura_tipi}.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text=f"✅ {fatura_tipi}.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_ve_isle_logo(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Logo XML klasörünü seçin")
    if secilen_klasor:
        xml_logo_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "logo_faturalar")

# ----------------- Uyumsoft XML -----------------

def xml_uyumsoft_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    global SON_DOSYA_YOLU
    veriler = []

    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()

            ns = XML_NS

            # Tarih ve Evrak No
            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)

            # VKN kontrolü - eğer Satıcı VKN'si hedef VKN ise KDV sıfırlanacak
            target_vkn = root.findtext(".//cac:AccountingSupplierParty//cbc:ID[@schemeID='VKN']", namespaces=ns)
            force_kdv_zero = (target_vkn == TARGET_VKN_FOR_ZERO_VAT)

            # Satıcı VKN ve Adı
            supplier_party = root.find(".//cac:AccountingSupplierParty/cac:Party", namespaces=ns)
            carikodu = supplier_party.findtext("cac:PartyIdentification/cbc:ID", namespaces=ns)
            cariadi = cari_adi_al(supplier_party, ns)

            # Yardımcı fonksiyon: güvenli float dönüşümü
            def _to_float(x):
                return float(str(x).replace(",", ".")) if x is not None else 0.0

            # Belge seviyesi AllowanceCharge kontrolü
            document_level_discount = 0.0
            doc_allowance = root.find("cac:AllowanceCharge", namespaces=ns)
            if doc_allowance is not None:
                charge_indicator = (doc_allowance.findtext("cbc:ChargeIndicator", namespaces=ns) or "").strip().lower()
                if charge_indicator == "false":
                    document_level_discount = _to_float(doc_allowance.findtext("cbc:Amount", namespaces=ns))

            # KDV'ye göre gruplama
            kdv_dict = {}
            lines = root.findall(".//cac:InvoiceLine", namespaces=ns)

            for line in lines:
                # BaseAmount (önce AllowanceCharge.BaseAmount'u dene, yoksa LineExtensionAmount)
                base_amount = 0.0
                ac_nodes = line.findall("cac:AllowanceCharge", namespaces=ns)
                if ac_nodes:
                    for ac in ac_nodes:
                        ba = ac.findtext("cbc:BaseAmount", namespaces=ns)
                        if ba is not None:
                            base_amount = _to_float(ba)
                            break
                if base_amount == 0.0:
                    base_amount = _to_float(line.findtext("cbc:LineExtensionAmount", namespaces=ns))

                # İndirim toplamı: ChargeIndicator == "false" olanların Amount toplamı
                discount_total = 0.0
                for ac in ac_nodes:
                    ch = (ac.findtext("cbc:ChargeIndicator", namespaces=ns) or "").strip().lower()
                    if ch == "false":
                        discount_total += _to_float(ac.findtext("cbc:Amount", namespaces=ns))

                # Satır değeri hesapla (KDV hariç)
                line_value = base_amount - discount_total
                if line_value < 0:
                    line_value = 0.0

                # KDV oranı
                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except (ValueError, TypeError):
                    kdv_orani = 0

                # VKN'ye göre KDV sıfırlama
                if force_kdv_zero:
                    kdv_orani = 0

                # Grupla (totalleri ve iskonto listelerini tut)
                if kdv_orani not in kdv_dict:
                    kdv_dict[kdv_orani] = {"total": 0.0, "line_discounts": []}
                kdv_dict[kdv_orani]["total"] += line_value
                kdv_dict[kdv_orani]["line_discounts"].append(discount_total)

            # Excel için veri satırlarını oluştur
            for kdv, grp_data in kdv_dict.items():
                toplam = grp_data["total"]
                
                # İskonto hesapla: Satır bazlı varsa onu kullan, yoksa belge bazlı kontrol et
                group_discount = sum(grp_data["line_discounts"])
                indtl = ""
                
                if group_discount > 0:
                    # Satır bazlı iskonto
                    indtl = round(group_discount, 2)
                elif document_level_discount > 0 and len(kdv_dict) == 1:
                    # Belge bazlı iskonto (sadece tek KDV oranı varsa)
                    indtl = round(document_level_discount, 2)
                
                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Alış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": str(kdv),    # KDV oranı = StokKodu
                    "STOKADI": "",           # Her zaman boş
                    "BIRIM": "ADET",         # Sabit
                    "MIKTAR": 1,             # Sabit
                    "DOVIZC": "TL",          # Sabit
                    "BIRIMFIYAT": round(toplam, 2),
                    "INDTL": indtl if indtl else "",
                    "KDVY": kdv,
                    "TUTARTL": round(toplam, 2),
                    "FATURACK": "",
                })

        except Exception as e:
            label_sonuc.configure(
                text=f"HATA: {os.path.basename(xml_dosya)} - {e}",
                text_color="red"
            )

    # Kaydetme işlemi
    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, f"{fatura_tipi}.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(
            text=f"✅ {fatura_tipi}.xlsx masaüstüne kaydedildi.",
            text_color="green"
        )
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(
            text="Geçerli XML dosyası bulunamadı.",
            text_color="red"
        )

def klasor_sec_ve_isle_uyumsoft(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Uyumsoft XML klasörünü seçin")
    if secilen_klasor:
        xml_uyumsoft_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "uyumsoft_faturalar")

# ----------------- Müstahsil XML -----------------

def temizle_kelime(kelime):
    kelime = kelime.strip().lower()
    kelime = kelime.replace("ı", "i").replace("ş", "s").replace("ç", "c").replace("ğ", "g").replace("ü", "u").replace("ö", "o")
    kelime = unicodedata.normalize("NFD", kelime).encode("ascii", "ignore").decode("utf-8")
    return kelime

def xml_mustahsil_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    global SON_DOSYA_YOLU

    urun_kodlari = {
        "patates": "08120101",
        "kuru sogan": "08110101",
        "sarimsak": "08190101",
        "limon": "16030301"
    }

    ns = XML_NS

    veriler = []

    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()

            belge_no = root.findtext(".//cbc:ID", namespaces=ns)
            belge_tarih = root.findtext(".//cbc:IssueDate", namespaces=ns)

            # TC kimlik no
            tc_no = ""
            for elem in root.findall(".//cac:AccountingCustomerParty//cbc:ID", namespaces=ns):
                if elem.attrib.get("schemeID") == "TCKN":
                    tc_no = elem.text.strip()
                    break

            for line in root.findall(".//cac:CreditNoteLine", namespaces=ns):
                miktar = line.findtext(".//cbc:CreditedQuantity", namespaces=ns)
                birim_fiyat = line.findtext(".//cbc:PriceAmount", namespaces=ns)
                tutar = line.findtext(".//cbc:LineExtensionAmount", namespaces=ns)
                urun_adi_raw = line.findtext(".//cac:Item//cbc:Name", namespaces=ns)

                if not urun_adi_raw:
                    continue

                urun_adi = temizle_kelime(urun_adi_raw)
                urun_kodu = urun_kodlari.get(urun_adi, "")

                # Mahsül yılı: belge tarihinin yılından 1 çıkar (önceki yılın mahsulü)
                mahsul_yili = str(int(belge_tarih[:4]) - 1) if belge_tarih else ""
                
                veriler.append({
                    "kimlik tipi": "k",
                    "TC KİMLİK NO": tc_no,
                    "MİKTAR": int(float(miktar)) if miktar else "",
                    "BİRİM FİYATI": birim_fiyat.replace(".", ",") if birim_fiyat else "",
                    "SATIS_TUTAR": tutar.replace(".", ",") if tutar else "",
                    "BELGE NO": belge_no,
                    "BELGE TARİHİ": f"{int(belge_tarih[8:10])}/{int(belge_tarih[5:7])}/{belge_tarih[:4]}" if belge_tarih else "",
                    "ÜRÜN KODU": urun_kodu,
                    "MENŞE": "001",
                    "MAHSÜL YILI": mahsul_yili,
                    "SATIŞ ŞEKLİ": "011",
                    "BAĞKUR MİKTAR": "",
                    "YRD_MIKTAR": "",
                    "YRD_OLCU_BIRIM": ""
                })

        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")

    if veriler:
        df = pd.DataFrame(veriler)
        dosya_yolu = os.path.join(MASAUSTU, f"{fatura_tipi}.xlsx")
        df.to_excel(dosya_yolu, index=False)
        label_sonuc.configure(text=f"✅ {fatura_tipi}.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_ve_isle_mustahsil(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Müstahsil Borsa XML klasörünü seçin")
    if secilen_klasor:
        xml_mustahsil_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "mustahsil_listesi")

# ----------------- Alış XML -----------------

def xml_alis_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    global SON_DOSYA_YOLU
    veriler = []

    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            ns = XML_NS

            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)

            # Alış faturası: Cari = Satıcı (Supplier)
            supplier_party = root.find(".//cac:AccountingSupplierParty/cac:Party", namespaces=ns)
            carikodu = supplier_party.findtext("cac:PartyIdentification/cbc:ID", namespaces=ns) if supplier_party is not None else ""
            cariadi = cari_adi_al(supplier_party, ns) if supplier_party is not None else "Bilinmeyen Cari"

            doviz = root.findtext(".//cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"

            for line in root.findall(".//cac:InvoiceLine", namespaces=ns):
                miktar = int(float(line.findtext("cbc:InvoicedQuantity", default="0", namespaces=ns)))
                birim = "ADET"
                fiyat = float(line.findtext("cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext("cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext("cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns) or ""

                iskonto_element = line.find("cac:AllowanceCharge/cbc:Amount", namespaces=ns)
                iskonto = float(iskonto_element.text.replace(",", ".")) if iskonto_element is not None else 0.0

                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except (ValueError, TypeError):
                    kdv_orani = 0

                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Alış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": birim,
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto,
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat * miktar,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")

    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, f"{fatura_tipi}.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text=f"✅ {fatura_tipi}.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_ve_isle_alis(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Alış XML klasörünü seçin")
    if secilen_klasor:
        xml_alis_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "mustahsil_alis")

# ----------------- Müstahsil XML (Makbuz) -----------------

def xml_mustahsil_makbuz_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    global SON_DOSYA_YOLU
    veriler = []

    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            ns = XML_NS

            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)

            # Müstahsil makbuzu: Cari = Müstahsil (AccountingCustomerParty)
            customer_party = root.find(".//cac:AccountingCustomerParty/cac:Party", namespaces=ns)
            carikodu = ""
            if customer_party is not None:
                for id_elem in customer_party.findall("cac:PartyIdentification/cbc:ID", namespaces=ns):
                    carikodu = id_elem.text.strip() if id_elem.text else ""
                    break
            cariadi = cari_adi_al(customer_party, ns) if customer_party is not None else "Bilinmeyen Cari"

            doviz = root.findtext(".//cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"

            for line in root.findall(".//cac:CreditNoteLine", namespaces=ns):
                # Miktar ve birim bilgisi
                miktar_element = line.find(".//cbc:CreditedQuantity", namespaces=ns)
                if miktar_element is not None:
                    miktar = int(float(miktar_element.text))
                    # Birim kodunu unitCode attribute'undan al
                    birim_code = miktar_element.attrib.get("unitCode", "KGM")
                    # KGM → KG, diğerleri olduğu gibi
                    birim = "KG" if birim_code == "KGM" else birim_code
                else:
                    miktar = 0
                    birim = "KG"
                
                fiyat = float(line.findtext(".//cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext(".//cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext(".//cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns) or ""

                iskonto_element = line.find(".//cac:AllowanceCharge/cbc:Amount", namespaces=ns)
                iskonto = float(iskonto_element.text.replace(",", ".")) if iskonto_element is not None else 0.0

                # Müstahsilde KDV hep 0
                kdv_orani = 0

                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Müstahsil Makbuzu",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": birim,
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto,
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat * miktar,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")

    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, f"{fatura_tipi}.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text=f"✅ {fatura_tipi}.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_ve_isle_mustahsil_xml(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Müstahsil XML klasörünü seçin")
    if secilen_klasor:
        xml_mustahsil_makbuz_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "mustahsil_makbuz")

# ----------------- Hepsiburada XML -----------------

def xml_hepsiburada_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    """
    Hepsiburada XML faturalarını okur ve Excel'e yazar.
    
    Args:
        fatura_tipi: "alis" veya "satis"
    """
    global SON_DOSYA_YOLU
    
    # Fatura tipine göre ayarlar
    if fatura_tipi == "alis":
        evrak_turu = "Alış Faturası"
        party_path = ".//cac:AccountingSupplierParty/cac:Party"
        dosya_adi = "hepsiburada_alis.xlsx"
    else:
        evrak_turu = "Satış Faturası"
        party_path = ".//cac:AccountingCustomerParty/cac:Party"
        dosya_adi = "hepsiburada_satis.xlsx"
    
    veriler = []
    ns = XML_NS
    
    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            
            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)
            
            # Cari bilgisi (fatura tipine göre)
            party = root.find(party_path, namespaces=ns)
            carikodu = party.findtext("cac:PartyIdentification/cbc:ID", namespaces=ns) if party is not None else ""
            cariadi = cari_adi_al(party, ns) if party is not None else "Bilinmeyen Cari"
            
            doviz = root.findtext(".//cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"
            
            for line in root.findall(".//cac:InvoiceLine", namespaces=ns):
                miktar = int(float(line.findtext("cbc:InvoicedQuantity", default="0", namespaces=ns)))
                fiyat = float(line.findtext("cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext("cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext("cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns) or ""
                
                # İskonto
                iskonto_element = line.find("cac:AllowanceCharge/cbc:Amount", namespaces=ns)
                iskonto = float(iskonto_element.text.replace(",", ".")) if iskonto_element is not None else 0.0
                
                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except (ValueError, TypeError):
                    kdv_orani = 0
                
                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": evrak_turu,
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": "ADET",
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto if iskonto else "",
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat * miktar,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")
    
    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, dosya_adi)
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text=f"✅ {dosya_adi} masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_hepsiburada_alis(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Hepsiburada Alış Faturası klasörünü seçin")
    if secilen_klasor:
        xml_hepsiburada_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "alis")

def klasor_sec_hepsiburada_satis(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Hepsiburada Satış Faturası klasörünü seçin")
    if secilen_klasor:
        xml_hepsiburada_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "satis")

# ----------------- Trendyol XML -----------------

def xml_trendyol_satis_oku_ve_yaz(klasor, label_sonuc, buton_ac):
    global SON_DOSYA_YOLU
    veriler = []
    ns = XML_NS
    
    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            
            tarih_raw = root.findtext("cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext("cbc:ID", namespaces=ns)
            
            # Satış faturası: Cari = Alıcı (Customer)
            customer_party = root.find("cac:AccountingCustomerParty/cac:Party", namespaces=ns)
            carikodu = customer_party.findtext("cac:PartyIdentification/cbc:ID", namespaces=ns) if customer_party is not None else ""
            cariadi = cari_adi_al(customer_party, ns) if customer_party is not None else "Bilinmeyen Cari"
            
            doviz = root.findtext("cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"
            
            # Fatura düzeyinde toplam iskonto
            toplam_iskonto_element = root.findtext("cac:LegalMonetaryTotal/cbc:AllowanceTotalAmount", namespaces=ns)
            toplam_iskonto = float(toplam_iskonto_element.replace(",", ".")) if toplam_iskonto_element else 0.0
            
            lines = root.findall("cac:InvoiceLine", namespaces=ns)
            for idx, line in enumerate(lines):
                miktar = int(float(line.findtext("cbc:InvoicedQuantity", default="0", namespaces=ns)))
                fiyat = float(line.findtext("cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext("cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext("cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns) or ""
                
                # İskonto: toplam iskontoyu sadece ilk satıra yaz, yoksa boş
                iskonto = toplam_iskonto if idx == 0 and toplam_iskonto else ""
                
                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except (ValueError, TypeError):
                    kdv_orani = 0
                
                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Satış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": "ADET",
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto,
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat * miktar,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")
    
    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, "trendyol_satis.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text="✅ trendyol_satis.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_trendyol_alis(label_sonuc, buton_ac):
    # Henüz Trendyol Alış faturası desteği yok
    label_sonuc.configure(text="⚠️ Trendyol Alış faturası henüz desteklenmiyor.", text_color="orange")

def klasor_sec_trendyol_satis(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Trendyol Satış Faturası klasörünü seçin")
    if secilen_klasor:
        xml_trendyol_satis_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac)

# ----------------- GUI -----------------

app = ctk.CTk()
app.title("Fatura Yönetim Sistemi")
app.geometry("650x610")
app.resizable(False, False)

# Açık popup menüyü takip etmek için
aktif_menu = None

# ----------------- Sayfa Yönetimi -----------------

def show_welcome():
    """Ana karşılama sayfasını göster"""
    xml_frame.pack_forget()
    kontrol_frame.pack_forget()
    welcome_frame.pack(expand=True, fill="both", padx=40, pady=30)

def show_xml():
    """XML aktarma sayfasını göster"""
    welcome_frame.pack_forget()
    kontrol_frame.pack_forget()
    xml_frame.pack(expand=True, fill="both", padx=40, pady=30)

def show_kontrol():
    """Kontrol sayfasını göster"""
    welcome_frame.pack_forget()
    xml_frame.pack_forget()
    kontrol_frame.pack(expand=True, fill="both", padx=40, pady=30)

# ----------------- Popup Menü Fonksiyonları -----------------

popup_aktif = False  # Popup durumunu takip et

def popup_menu_kapat():
    global aktif_menu, popup_aktif
    popup_aktif = False
    if aktif_menu is not None:
        try:
            aktif_menu.destroy()
        except Exception:
            # Pencere zaten kapatılmış olabilir
            pass
        aktif_menu = None

def on_app_click(event):
    """Ana pencereye tıklanınca popup'ı kapat"""
    global popup_aktif
    if popup_aktif and aktif_menu is not None:
        # Popup içine tıklanmadıysa kapat
        widget = event.widget
        # Popup'ın child'ı değilse kapat
        if str(widget).startswith(".!ctktoplevel"):
            return  # Popup içi, kapatma
        popup_menu_kapat()

def show_popup_menu(button, label_sonuc, buton_ac, alis_func, satis_func):
    """Genel popup menü fonksiyonu - Alış/Satış seçenekleri gösterir"""
    global aktif_menu, popup_aktif
    popup_menu_kapat()
    
    x = button.winfo_rootx()
    y = button.winfo_rooty() + button.winfo_height() + 5
    
    aktif_menu = ctk.CTkToplevel(app)
    aktif_menu.geometry(f"160x90+{x}+{y}")
    aktif_menu.overrideredirect(True)
    aktif_menu.attributes("-topmost", True)
    aktif_menu.configure(fg_color="#3b3b3b")
    
    ctk.CTkButton(aktif_menu, text="📥 Alış Faturası", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), alis_func(label_sonuc, buton_ac)]).pack(pady=(8, 4), padx=5)
    ctk.CTkButton(aktif_menu, text="📤 Satış Faturası", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), satis_func(label_sonuc, buton_ac)]).pack(pady=(0, 8), padx=5)
    
    # Popup aktif olarak işaretle (biraz gecikmeyle)
    def activate_popup():
        global popup_aktif
        popup_aktif = True
    app.after(150, activate_popup)
    
    # Escape tuşu ile kapatma
    aktif_menu.bind("<Escape>", lambda e: popup_menu_kapat())

def show_hepsiburada_menu(button, label_sonuc, buton_ac):
    show_popup_menu(button, label_sonuc, buton_ac, klasor_sec_hepsiburada_alis, klasor_sec_hepsiburada_satis)

def show_trendyol_menu(button, label_sonuc, buton_ac):
    show_popup_menu(button, label_sonuc, buton_ac, klasor_sec_trendyol_alis, klasor_sec_trendyol_satis)

def show_hizlibilisim_xml_menu(button, label_sonuc, buton_ac):
    """Hızlıbilişim XML için alış/satış menüsü - Hepsiburada ile aynı fonksiyonları kullanır"""
    show_popup_menu(button, label_sonuc, buton_ac, klasor_sec_hepsiburada_alis, klasor_sec_hepsiburada_satis)

def show_mustahsil_menu(button, label_sonuc, buton_ac):
    """Müstahsil butonu için 3 seçenekli popup menü: Borsa, Alış XML, Müstahsil XML"""
    global aktif_menu, popup_aktif
    popup_menu_kapat()
    
    x = button.winfo_rootx()
    y = button.winfo_rooty() + button.winfo_height() + 5
    
    aktif_menu = ctk.CTkToplevel(app)
    aktif_menu.geometry(f"160x135+{x}+{y}")
    aktif_menu.overrideredirect(True)
    aktif_menu.attributes("-topmost", True)
    aktif_menu.configure(fg_color="#3b3b3b")
    
    ctk.CTkButton(aktif_menu, text="🥔 Müstahsil Borsa", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), klasor_sec_ve_isle_mustahsil(label_sonuc, buton_ac)]).pack(pady=(8, 4), padx=5)
    ctk.CTkButton(aktif_menu, text="📦 Alış XML", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), klasor_sec_ve_isle_alis(label_sonuc, buton_ac)]).pack(pady=(0, 4), padx=5)
    ctk.CTkButton(aktif_menu, text="📋 Müstahsil XML", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), klasor_sec_ve_isle_mustahsil_xml(label_sonuc, buton_ac)]).pack(pady=(0, 8), padx=5)
    
    # Popup aktif olarak işaretle (biraz gecikmeyle)
    def activate_popup():
        global popup_aktif
        popup_aktif = True
    app.after(150, activate_popup)
    
    # Escape tuşu ile kapatma
    aktif_menu.bind("<Escape>", lambda e: popup_menu_kapat())

def show_uyumsoft_xml_menu(button, label_sonuc, buton_ac):
    """Uyumsoft XML için ACE/Turkay dropdown menüsü"""
    global aktif_menu, popup_aktif
    popup_menu_kapat()
    
    x = button.winfo_rootx()
    y = button.winfo_rooty() + button.winfo_height() + 5
    
    aktif_menu = ctk.CTkToplevel(app)
    aktif_menu.geometry(f"160x90+{x}+{y}")
    aktif_menu.overrideredirect(True)
    aktif_menu.attributes("-topmost", True)
    aktif_menu.configure(fg_color="#3b3b3b")
    
    ctk.CTkButton(aktif_menu, text="🏢 ACE", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), show_info_popup("Bu özellik henüz aktif değil.")]).pack(pady=(8, 4), padx=5)
    ctk.CTkButton(aktif_menu, text="🏭 Turkay", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), klasor_sec_ve_isle_uyumsoft(label_sonuc, buton_ac)]).pack(pady=(0, 8), padx=5)
    
    # Popup aktif olarak işaretle (biraz gecikmeyle)
    def activate_popup():
        global popup_aktif
        popup_aktif = True
    app.after(150, activate_popup)
    
    # Escape tuşu ile kapatma
    aktif_menu.bind("<Escape>", lambda e: popup_menu_kapat())

# ================= KARŞILAMA SAYFASI =================

welcome_frame = ctk.CTkFrame(app, fg_color="transparent")

welcome_inner = ctk.CTkFrame(welcome_frame, fg_color="#2b2b2b", corner_radius=15)
welcome_inner.pack(expand=True, fill="both")

# Başlık
welcome_title = ctk.CTkLabel(
    welcome_inner, 
    text="📄 Fatura Yönetim Sistemi", 
    font=ctk.CTkFont(size=28, weight="bold")
)
welcome_title.pack(pady=(30, 8))

welcome_desc = ctk.CTkLabel(
    welcome_inner,
    text="Yapmak istediğiniz işlemi seçin",
    font=ctk.CTkFont(size=16),
    text_color="#a0a0a0"
)
welcome_desc.pack(pady=(0, 20))

# Kart Container - tüm alanı kaplasın
cards_frame = ctk.CTkFrame(welcome_inner, fg_color="transparent")
cards_frame.pack(expand=True, fill="both", padx=25, pady=(0, 25))
cards_frame.grid_columnconfigure(0, weight=1)
cards_frame.grid_columnconfigure(1, weight=1)
cards_frame.grid_rowconfigure(0, weight=1)

# Kart hover efekti ve tıklama için yardımcı fonksiyon
def make_card_clickable(card, all_widgets, command):
    """Karta hover efekti ve tıklama özelliği ekler"""
    normal_color = "#363636"
    hover_color = "#454545"
    click_color = "#505050"
    
    def on_enter(e):
        card.configure(fg_color=hover_color)
    
    def on_leave(e):
        card.configure(fg_color=normal_color)
    
    def on_click(e):
        card.configure(fg_color=click_color)
        card.after(100, lambda: [card.configure(fg_color=normal_color), command()])
    
    # Kart ve tüm child widget'lara event bağla
    for widget in all_widgets:
        widget.bind("<Enter>", on_enter)
        widget.bind("<Leave>", on_leave)
        widget.bind("<Button-1>", on_click)
        widget.configure(cursor="hand2")

# XML Kartı
xml_card = ctk.CTkFrame(cards_frame, fg_color="#363636", corner_radius=12)
xml_card.grid(row=0, column=0, padx=(0, 10), sticky="nsew")

xml_card_content = ctk.CTkFrame(xml_card, fg_color="transparent")
xml_card_content.pack(expand=True)

xml_icon = ctk.CTkLabel(xml_card_content, text="📁", font=ctk.CTkFont(size=56))
xml_icon.pack(pady=(0, 15))

xml_card_title = ctk.CTkLabel(xml_card_content, text="XML Aktarma", font=ctk.CTkFont(size=22, weight="bold"))
xml_card_title.pack(pady=(0, 10))

xml_card_desc = ctk.CTkLabel(
    xml_card_content, 
    text="XML faturalarını\nExcel'e aktarın",
    font=ctk.CTkFont(size=14),
    text_color="#b0b0b0",
    justify="center"
)
xml_card_desc.pack()

# XML kartını tıklanabilir yap
make_card_clickable(xml_card, [xml_card, xml_card_content, xml_icon, xml_card_title, xml_card_desc], show_xml)

# Kontrol Kartı
kontrol_card = ctk.CTkFrame(cards_frame, fg_color="#363636", corner_radius=12)
kontrol_card.grid(row=0, column=1, padx=(10, 0), sticky="nsew")

kontrol_card_content = ctk.CTkFrame(kontrol_card, fg_color="transparent")
kontrol_card_content.pack(expand=True)

kontrol_icon = ctk.CTkLabel(kontrol_card_content, text="🔍", font=ctk.CTkFont(size=56))
kontrol_icon.pack(pady=(0, 15))

kontrol_card_title = ctk.CTkLabel(kontrol_card_content, text="Fatura Kontrol", font=ctk.CTkFont(size=22, weight="bold"))
kontrol_card_title.pack(pady=(0, 10))

kontrol_card_desc = ctk.CTkLabel(
    kontrol_card_content, 
    text="Zirve ve portal\nfaturalarını karşılaştırın",
    font=ctk.CTkFont(size=14),
    text_color="#b0b0b0",
    justify="center"
)
kontrol_card_desc.pack()

# Kontrol kartını tıklanabilir yap
make_card_clickable(kontrol_card, [kontrol_card, kontrol_card_content, kontrol_icon, kontrol_card_title, kontrol_card_desc], show_kontrol)

# ================= XML SAYFASI =================

xml_frame = ctk.CTkFrame(app, fg_color="transparent")

xml_inner = ctk.CTkFrame(xml_frame, fg_color="#2b2b2b", corner_radius=15)
xml_inner.pack(expand=True, fill="both")

# Geri Butonu
xml_back_frame = ctk.CTkFrame(xml_inner, fg_color="transparent")
xml_back_frame.pack(fill="x", padx=15, pady=(15, 0))

xml_back_btn = ctk.CTkButton(
    xml_back_frame, 
    text="← Geri", 
    width=80, 
    height=30,
    fg_color="transparent",
    border_width=1,
    border_color="#555555",
    hover_color="#404040",
    command=show_welcome
)
xml_back_btn.pack(side="left")

# Başlık
xml_title = ctk.CTkLabel(xml_inner, text="📄 XML Aktarma", font=ctk.CTkFont(size=24, weight="bold"))
xml_title.pack(pady=(10, 5))

xml_aciklama = ctk.CTkLabel(
    xml_inner,
    text="XML formatındaki faturaları seçerek Excel çıktısı oluşturabilirsiniz.",
    font=ctk.CTkFont(size=14),
    text_color="#d0d0d0",
    justify="center"
)
xml_aciklama.pack(pady=(0, 20))

# Satır 1: Logo ve Uyumsoft Butonları
xml_row1 = ctk.CTkFrame(xml_inner, fg_color="transparent")
xml_row1.pack(pady=8)

# Logo
logo_col = ctk.CTkFrame(xml_row1, fg_color="transparent")
logo_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(logo_col, text="Logo", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(logo_col, text="📂 Aktar", width=160, height=40,
              command=lambda: klasor_sec_ve_isle_logo(xml_label_sonuc, xml_buton_ac)).pack()

# Uyumsoft
uyumsoft_col = ctk.CTkFrame(xml_row1, fg_color="transparent")
uyumsoft_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(uyumsoft_col, text="Uyumsoft", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
uyumsoft_xml_buton = ctk.CTkButton(uyumsoft_col, text="🧾 Aktar", width=160, height=40)
uyumsoft_xml_buton.pack()

# Satır 2: Müstahsil ve Hızlıbilişim
xml_row2 = ctk.CTkFrame(xml_inner, fg_color="transparent")
xml_row2.pack(pady=8)

# Müstahsil
mustahsil_col = ctk.CTkFrame(xml_row2, fg_color="transparent")
mustahsil_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(mustahsil_col, text="Müstahsil", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
mustahsil_buton = ctk.CTkButton(mustahsil_col, text="🥔 Aktar", width=160, height=40)
mustahsil_buton.pack()

# Hızlıbilişim
hizlibilisim_xml_col = ctk.CTkFrame(xml_row2, fg_color="transparent")
hizlibilisim_xml_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(hizlibilisim_xml_col, text="Hızlıbilişim", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
hizlibilisim_xml_buton = ctk.CTkButton(hizlibilisim_xml_col, text="⚡ Aktar", width=160, height=40)
hizlibilisim_xml_buton.pack()

# Satır 3: Hepsiburada ve Trendyol
xml_row3 = ctk.CTkFrame(xml_inner, fg_color="transparent")
xml_row3.pack(pady=8)

# Hepsiburada
hb_col = ctk.CTkFrame(xml_row3, fg_color="transparent")
hb_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(hb_col, text="Hepsiburada", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
hb_buton = ctk.CTkButton(hb_col, text="🛒 Aktar", width=160, height=40)
hb_buton.pack()

# Trendyol
ty_col = ctk.CTkFrame(xml_row3, fg_color="transparent")
ty_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(ty_col, text="Trendyol", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ty_buton = ctk.CTkButton(ty_col, text="🛍️ Aktar", width=160, height=40)
ty_buton.pack()

# Sonuç mesajı
xml_label_sonuc = ctk.CTkLabel(xml_inner, text="", font=ctk.CTkFont(size=16))
xml_label_sonuc.pack(pady=10)

# Excel aç butonu
xml_buton_ac = ctk.CTkButton(xml_inner, text="📂 Oluşturulan Excel Dosyasını Aç", command=dosyayi_ac, state="disabled", width=280, height=40)
xml_buton_ac.pack(pady=(10, 30))

# Popup menü komutlarını bağla (label_sonuc ve buton_ac tanımlandıktan sonra)
uyumsoft_xml_buton.configure(command=lambda: show_uyumsoft_xml_menu(uyumsoft_xml_buton, xml_label_sonuc, xml_buton_ac))
mustahsil_buton.configure(command=lambda: show_mustahsil_menu(mustahsil_buton, xml_label_sonuc, xml_buton_ac))
hb_buton.configure(command=lambda: show_hepsiburada_menu(hb_buton, xml_label_sonuc, xml_buton_ac))
ty_buton.configure(command=lambda: show_trendyol_menu(ty_buton, xml_label_sonuc, xml_buton_ac))
hizlibilisim_xml_buton.configure(command=lambda: show_hizlibilisim_xml_menu(hizlibilisim_xml_buton, xml_label_sonuc, xml_buton_ac))

# ================= KONTROL SAYFASI =================

kontrol_frame = ctk.CTkFrame(app, fg_color="transparent")

kontrol_inner = ctk.CTkFrame(kontrol_frame, fg_color="#2b2b2b", corner_radius=15)
kontrol_inner.pack(expand=False, fill="x", padx=40, pady=30)

# Geri Butonu
kontrol_back_frame = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_back_frame.pack(fill="x", padx=15, pady=(15, 10))

kontrol_back_btn = ctk.CTkButton(
    kontrol_back_frame, 
    text="← Geri", 
    width=80, 
    height=30,
    fg_color="transparent",
    border_width=1,
    border_color="#555555",
    hover_color="#404040",
    command=show_welcome
)
kontrol_back_btn.pack(side="left")

# Başlık
kontrol_title = ctk.CTkLabel(kontrol_inner, text="🔍 Fatura Kontrol", font=ctk.CTkFont(size=24, weight="bold"))
kontrol_title.pack(pady=(5, 2))

kontrol_aciklama = ctk.CTkLabel(
    kontrol_inner,
    text="Zirve ve portal faturalarını karşılaştırarak\neksik veya hatalı girişleri tespit edin.",
    font=ctk.CTkFont(size=14),
    text_color="#d0d0d0",
    justify="center"
)
kontrol_aciklama.pack(pady=(2, 25))

# Kontrol popup fonksiyonu
def show_kontrol_popup(portal_adi, kontrol_func=None):
    """Excel karşılaştırma popup'ı göster"""
    popup = ctk.CTkToplevel(app)
    popup.title(f"{portal_adi} Kontrol")
    popup.geometry("500x360")
    popup.resizable(False, False)
    popup.attributes("-topmost", True)
    popup.grab_set()  # Modal yap
    
    # Pencereyi ortala
    popup.update_idletasks()
    x = app.winfo_x() + (app.winfo_width() // 2) - 250
    y = app.winfo_y() + (app.winfo_height() // 2) - 180
    popup.geometry(f"500x360+{x}+{y}")
    
    # Seçilen dosyaları takip et
    secilen_dosyalar = {"zirve": None, "portal": []}
    
    # Başlık
    ctk.CTkLabel(popup, text=f"🔍 {portal_adi} Fatura Kontrol", 
                 font=ctk.CTkFont(size=20, weight="bold")).pack(pady=(20, 5))
    ctk.CTkLabel(popup, text="Karşılaştırmak için Excel dosyası/dosyaları seçin", 
                 font=ctk.CTkFont(size=12), text_color="#a0a0a0").pack(pady=(0, 20))
    
    # Dosya seçim alanları
    files_frame = ctk.CTkFrame(popup, fg_color="transparent")
    files_frame.pack(fill="x", padx=30)
    
    # Sol: Zirve Excel
    zirve_frame = ctk.CTkFrame(files_frame, fg_color="#363636", corner_radius=10)
    zirve_frame.grid(row=0, column=0, padx=(0, 10), sticky="nsew")
    files_frame.grid_columnconfigure(0, weight=1)
    
    ctk.CTkLabel(zirve_frame, text="Zirve Excel", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 5))
    zirve_status = ctk.CTkLabel(zirve_frame, text="📄 Dosya seçilmedi", font=ctk.CTkFont(size=11), text_color="#808080")
    zirve_status.pack(pady=(0, 10))
    
    def select_zirve():
        # Popup'ı gizle
        popup.withdraw()
        popup.update_idletasks()
        
        dosya = filedialog.askopenfilename(
            title="Zirve Excel Dosyası Seçin",
            filetypes=[("Excel Dosyaları", "*.xlsx *.xls")]
        )
        
        # Popup'ı tekrar göster
        popup.deiconify()
        popup.lift()
        popup.focus_force()
        
        if dosya:
            secilen_dosyalar["zirve"] = dosya
            dosya_adi = os.path.basename(dosya)
            if len(dosya_adi) > 20:
                dosya_adi = dosya_adi[:17] + "..."
            zirve_status.configure(text=f"✅ {dosya_adi}", text_color="#4CAF50")
            check_ready()
    
    ctk.CTkButton(zirve_frame, text="📂 Seç", width=100, height=32, command=select_zirve).pack(pady=(0, 15))
    
    # Sağ: Portal Excel (Çoklu Seçim)
    portal_frame = ctk.CTkFrame(files_frame, fg_color="#363636", corner_radius=10)
    portal_frame.grid(row=0, column=1, padx=(10, 0), sticky="nsew")
    files_frame.grid_columnconfigure(1, weight=1)
    
    ctk.CTkLabel(portal_frame, text=f"{portal_adi} Excel", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 5))
    portal_status = ctk.CTkLabel(portal_frame, text="📄 Dosya seçilmedi", font=ctk.CTkFont(size=11), text_color="#808080")
    portal_status.pack(pady=(0, 10))
    
    def select_portal():
        # Popup'ı gizle
        popup.withdraw()
        popup.update_idletasks()
        
        dosyalar = filedialog.askopenfilenames(
            title=f"{portal_adi} Excel Dosyası(ları) Seçin (Çoklu seçim için Ctrl tuşu)",
            filetypes=[("Excel Dosyaları", "*.xlsx *.xls")]
        )
        
        # Popup'ı tekrar göster
        popup.deiconify()
        popup.lift()
        popup.focus_force()
        
        if dosyalar:
            secilen_dosyalar["portal"] = list(dosyalar)
            dosya_sayisi = len(dosyalar)
            
            if dosya_sayisi == 1:
                dosya_adi = os.path.basename(dosyalar[0])
                if len(dosya_adi) > 18:
                    dosya_adi = dosya_adi[:15] + "..."
                portal_status.configure(text=f"✅ {dosya_adi}", text_color="#4CAF50")
            else:
                portal_status.configure(text=f"✅ {dosya_sayisi} dosya seçildi", text_color="#4CAF50")
            
            check_ready()
    
    ctk.CTkButton(portal_frame, text="📂 Dosyalar Seç", width=110, height=32, command=select_portal).pack(pady=(0, 15))
    
    # Kontrol Et butonu (başlangıçta devre dışı)
    def run_kontrol():
        if secilen_dosyalar["zirve"] and secilen_dosyalar["portal"]:
            popup.destroy()
            if kontrol_func:
                kontrol_func(secilen_dosyalar["zirve"], secilen_dosyalar["portal"])
            else:
                # Henüz implementasyon yok
                show_info_popup("Bu özellik henüz aktif değil.")
    
    kontrol_btn = ctk.CTkButton(popup, text="🔍 Kontrol Et", width=200, height=40, 
                                 command=run_kontrol, state="disabled",
                                 fg_color="#555555", hover_color="#555555")
    kontrol_btn.pack(pady=(25, 15))
    
    def check_ready():
        # Zirve seçilmiş ve Portal'de en az 1 dosya seçilmişse
        if secilen_dosyalar["zirve"] and len(secilen_dosyalar["portal"]) > 0:
            kontrol_btn.configure(state="normal", fg_color="#1f6aa5", hover_color="#144870")
    
    # İptal butonu
    ctk.CTkButton(popup, text="İptal", width=100, height=30, 
                  fg_color="transparent", border_width=1, border_color="#555555",
                  hover_color="#404040", command=popup.destroy).pack()

def show_info_popup(mesaj):
    """Bilgi popup'ı göster - Dinamik boyutlandırma ve kopyalanabilir metin"""
    popup = ctk.CTkToplevel(app)
    popup.title("Bilgi")
    popup.attributes("-topmost", True)
    popup.grab_set()
    
    # Mesaj uzunluğuna göre dinamik boyut hesapla
    satir_sayisi = mesaj.count('\n') + 1
    karakter_sayisi = max(len(satir) for satir in mesaj.split('\n'))
    
    # Minimum ve maksimum boyutlar
    min_width, max_width = 350, 800
    min_height, max_height = 200, 600
    
    # Genişlik: karakter sayısına göre (ortalama 8 piksel/karakter)
    width = min(max(min_width, karakter_sayisi * 8 + 100), max_width)
    # Yükseklik: satır sayısına göre (ortalama 25 piksel/satır) + ekstra boşluklar
    height = min(max(min_height, satir_sayisi * 25 + 180), max_height)
    
    popup.geometry(f"{width}x{height}")
    popup.resizable(True, True)  # Kullanıcı manuel olarak da büyütebilir
    
    # Pencereyi ortala
    popup.update_idletasks()
    x = app.winfo_x() + (app.winfo_width() // 2) - (width // 2)
    y = app.winfo_y() + (app.winfo_height() // 2) - (height // 2)
    popup.geometry(f"{width}x{height}+{x}+{y}")
    
    # İkon
    ctk.CTkLabel(popup, text="⚠️", font=ctk.CTkFont(size=36)).pack(pady=(15, 5))
    
    # Kopyalanabilir metin kutusu
    text_frame = ctk.CTkFrame(popup, fg_color="transparent")
    text_frame.pack(fill="both", expand=True, padx=20, pady=(0, 15))
    
    textbox = ctk.CTkTextbox(text_frame, font=ctk.CTkFont(size=12), 
                             wrap="word", activate_scrollbars=True)
    textbox.pack(fill="both", expand=True)
    textbox.insert("1.0", mesaj)
    textbox.configure(state="normal")  # Seçilebilir ve kopyalanabilir
    
    # Buton çerçevesi
    btn_frame = ctk.CTkFrame(popup, fg_color="transparent")
    btn_frame.pack(pady=(0, 15))
    
    # Tümünü Kopyala butonu
    def kopyala():
        popup.clipboard_clear()
        popup.clipboard_append(mesaj)
        kopyala_btn.configure(text="✓ Kopyalandı", fg_color="#4CAF50")
        popup.after(1500, lambda: kopyala_btn.configure(text="📋 Tümünü Kopyala", fg_color="#1f6aa5"))
    
    kopyala_btn = ctk.CTkButton(btn_frame, text="📋 Tümünü Kopyala", width=140, height=32, 
                                command=kopyala, fg_color="#1f6aa5", hover_color="#144870")
    kopyala_btn.pack(side="left", padx=5)
    
    ctk.CTkButton(btn_frame, text="Tamam", width=100, height=32, command=popup.destroy).pack(side="left", padx=5)

# Kontrol callback fabrika fonksiyonu
def _kontrol_sonuc_popup_goster(karsilastir_func):
    """Kontrol sonucu popup'ı gösteren callback oluşturur"""
    def callback(zirve_path, portal_path):
        sonuc_dosyasi, mesaj = karsilastir_func(zirve_path, portal_path)
        
        if sonuc_dosyasi:
            popup = ctk.CTkToplevel(app)
            popup.title("Kontrol Sonucu")
            popup.geometry("420x320")
            popup.resizable(False, False)
            popup.attributes("-topmost", True)
            popup.grab_set()
            
            popup.update_idletasks()
            x = app.winfo_x() + (app.winfo_width() // 2) - 210
            y = app.winfo_y() + (app.winfo_height() // 2) - 160
            popup.geometry(f"420x320+{x}+{y}")
            
            ctk.CTkLabel(popup, text="🎉", font=ctk.CTkFont(size=36)).pack(pady=(20, 5))
            ctk.CTkLabel(popup, text=mesaj, font=ctk.CTkFont(size=13), 
                         justify="left").pack(pady=(0, 20), padx=20)
            
            btn_frame = ctk.CTkFrame(popup, fg_color="transparent")
            btn_frame.pack(pady=(0, 20))
            
            ctk.CTkButton(btn_frame, text="📂 Dosyayı Aç", width=150, height=35, 
                          command=lambda: [popup.destroy(), dosyayi_ac()]).pack(side="left", padx=5)
            ctk.CTkButton(btn_frame, text="Kapat", width=100, height=35, 
                          fg_color="transparent", border_width=1, border_color="#555555",
                          command=popup.destroy).pack(side="left", padx=5)
        else:
            show_info_popup(mesaj)
    return callback

# Kontrol callback'leri - fabrika fonksiyonuyla oluştur
# Uyumsoft için Zirve/Verim callback'leri
def uyumsoft_zirve_callback(zirve_path, portal_paths):
    return uyumsoft_zirve_karsilastir(zirve_path, portal_paths)

def uyumsoft_verim_callback(zirve_path, portal_paths):
    return uyumsoft_verim_karsilastir(zirve_path, portal_paths)

uyumsoft_zirve_kontrol_callback = _kontrol_sonuc_popup_goster(uyumsoft_zirve_callback)
uyumsoft_verim_kontrol_callback = _kontrol_sonuc_popup_goster(uyumsoft_verim_callback)
vega_kontrol_callback = _kontrol_sonuc_popup_goster(vega_karsilastir)
izibiz_kontrol_callback = _kontrol_sonuc_popup_goster(izibiz_karsilastir)

# Hızlıbilişim için Alış/Satış callback'leri
def hizlibilisim_alis_callback(zirve_path, portal_path):
    return hizlibilisim_karsilastir(zirve_path, portal_path, fatura_yonu='alis')

def hizlibilisim_satis_callback(zirve_path, portal_paths):
    return hizlibilisim_karsilastir(zirve_path, portal_paths, fatura_yonu='satis')

hizlibilisim_alis_kontrol_callback = _kontrol_sonuc_popup_goster(hizlibilisim_alis_callback)
hizlibilisim_satis_kontrol_callback = _kontrol_sonuc_popup_goster(hizlibilisim_satis_callback)

# Hepsiburada/Trendyol için callback (otomatik tespit)
hb_ve_trn_kontrol_callback = _kontrol_sonuc_popup_goster(hb_ve_trn_karsilastir)

# Henüz aktif olmayan özellikler için ortak fonksiyon
def henuz_aktif_degil():
    show_info_popup("Bu özellik henüz aktif değil.")

# Kontrol fonksiyonları
def kontrol_uyumsoft_zirve():
    show_kontrol_popup("Uyumsoft - Zirve", uyumsoft_zirve_kontrol_callback)

def show_uyumsoft_verim_popup():
    """Uyumsoft Verim kontrol popup'ı - Verim ve Uyumsoft dosyaları"""
    popup = ctk.CTkToplevel(app)
    popup.title("Uyumsoft - Verim Kontrol")
    popup.geometry("520x380")
    popup.resizable(False, False)
    popup.attributes("-topmost", True)
    popup.grab_set()
    
    popup.update_idletasks()
    x = app.winfo_x() + (app.winfo_width() // 2) - 260
    y = app.winfo_y() + (app.winfo_height() // 2) - 190
    popup.geometry(f"520x380+{x}+{y}")
    
    secilen_dosyalar = {"verim": [], "portal": []}
    
    ctk.CTkLabel(popup, text="🔍 Uyumsoft - Verim Fatura Kontrol", 
                 font=ctk.CTkFont(size=20, weight="bold")).pack(pady=(20, 5))
    ctk.CTkLabel(popup, text="Verim ve Uyumsoft Excel dosyalarını seçin", 
                 font=ctk.CTkFont(size=12), text_color="#a0a0a0").pack(pady=(0, 20))
    
    files_frame = ctk.CTkFrame(popup, fg_color="transparent")
    files_frame.pack(fill="x", padx=30)
    
    # Sol: Verim Excel
    verim_frame = ctk.CTkFrame(files_frame, fg_color="#363636", corner_radius=10)
    verim_frame.grid(row=0, column=0, padx=(0, 10), sticky="nsew")
    files_frame.grid_columnconfigure(0, weight=1)
    
    ctk.CTkLabel(verim_frame, text="Verim Excel", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 5))
    verim_status = ctk.CTkLabel(verim_frame, text="📄 Dosya seçilmedi", font=ctk.CTkFont(size=11), text_color="#808080")
    verim_status.pack(pady=(0, 10))
    
    def select_verim():
        popup.withdraw()
        popup.update_idletasks()
        dosyalar = filedialog.askopenfilenames(title="Verim Excel Dosyaları Seçin", filetypes=[("Excel Dosyaları", "*.xlsx *.xls")])
        popup.deiconify()
        popup.lift()
        popup.focus_force()
        if dosyalar:
            secilen_dosyalar["verim"] = list(dosyalar)
            dosya_sayisi = len(secilen_dosyalar["verim"])
            if dosya_sayisi == 1:
                dosya_adi = os.path.basename(dosyalar[0])
                if len(dosya_adi) > 18:
                    dosya_adi = dosya_adi[:15] + "..."
                verim_status.configure(text=f"✅ {dosya_adi}", text_color="#4CAF50")
            else:
                verim_status.configure(text=f"✅ {dosya_sayisi} dosya seçildi", text_color="#4CAF50")
            check_ready()
    
    ctk.CTkButton(verim_frame, text="📂 Dosyalar Seç", width=110, height=32, command=select_verim).pack(pady=(0, 15))
    
    # Sağ: Uyumsoft Excel
    portal_frame = ctk.CTkFrame(files_frame, fg_color="#363636", corner_radius=10)
    portal_frame.grid(row=0, column=1, padx=(10, 0), sticky="nsew")
    files_frame.grid_columnconfigure(1, weight=1)
    
    ctk.CTkLabel(portal_frame, text="Uyumsoft Excel", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 5))
    portal_status = ctk.CTkLabel(portal_frame, text="📄 Dosya seçilmedi", font=ctk.CTkFont(size=11), text_color="#808080")
    portal_status.pack(pady=(0, 10))
    
    def select_portal():
        popup.withdraw()
        popup.update_idletasks()
        dosyalar = filedialog.askopenfilenames(title="Uyumsoft Excel Dosyaları Seçin", filetypes=[("Excel Dosyaları", "*.xlsx *.xls")])
        popup.deiconify()
        popup.lift()
        popup.focus_force()
        if dosyalar:
            secilen_dosyalar["portal"] = list(dosyalar)
            dosya_sayisi = len(secilen_dosyalar["portal"])
            if dosya_sayisi == 1:
                dosya_adi = os.path.basename(dosyalar[0])
                if len(dosya_adi) > 18:
                    dosya_adi = dosya_adi[:15] + "..."
                portal_status.configure(text=f"✅ {dosya_adi}", text_color="#4CAF50")
            else:
                portal_status.configure(text=f"✅ {dosya_sayisi} dosya seçildi", text_color="#4CAF50")
            check_ready()
    
    ctk.CTkButton(portal_frame, text="📂 Dosyalar Seç", width=110, height=32, command=select_portal).pack(pady=(0, 15))
    
    def run_kontrol():
        if secilen_dosyalar["verim"] and secilen_dosyalar["portal"]:
            popup.destroy()
            uyumsoft_verim_kontrol_callback(secilen_dosyalar["verim"], secilen_dosyalar["portal"])
    
    kontrol_btn = ctk.CTkButton(popup, text="🔍 Kontrol Et", width=200, height=40, 
                                 command=run_kontrol, state="disabled",
                                 fg_color="#555555", hover_color="#555555")
    kontrol_btn.pack(pady=(25, 15))
    
    def check_ready():
        if secilen_dosyalar["verim"] and secilen_dosyalar["portal"]:
            kontrol_btn.configure(state="normal", fg_color="#1f6aa5", hover_color="#144870")
    
    ctk.CTkButton(popup, text="İptal", width=100, height=30, 
                  fg_color="transparent", border_width=1, border_color="#555555",
                  hover_color="#404040", command=popup.destroy).pack()

def kontrol_uyumsoft_verim():
    show_uyumsoft_verim_popup()

def kontrol_vega():
    show_kontrol_popup("Vega", vega_kontrol_callback)

def kontrol_izibiz():
    show_kontrol_popup("İzibiz", izibiz_kontrol_callback)

def kontrol_hb_ve_trn():
    show_kontrol_popup("Hepsiburada/Trendyol", hb_ve_trn_kontrol_callback)

# Hızlıbilişim için özel kontrol popup'ı (tek dosya seçimi)
def show_hizlibilisim_alis_popup():
    """Hızlıbilişim Alış kontrol popup'ı - tek dosya seçimi"""
    popup = ctk.CTkToplevel(app)
    popup.title("Hızlıbilişim Alış Kontrol")
    popup.geometry("500x360")
    popup.resizable(False, False)
    popup.attributes("-topmost", True)
    popup.grab_set()
    
    popup.update_idletasks()
    x = app.winfo_x() + (app.winfo_width() // 2) - 250
    y = app.winfo_y() + (app.winfo_height() // 2) - 180
    popup.geometry(f"500x360+{x}+{y}")
    
    secilen_dosyalar = {"zirve": None, "portal": None}
    
    ctk.CTkLabel(popup, text="🔍 Hızlıbilişim Alış Fatura Kontrol", 
                 font=ctk.CTkFont(size=20, weight="bold")).pack(pady=(20, 5))
    ctk.CTkLabel(popup, text="Gelen faturalar için karşılaştırma (tek dosya)", 
                 font=ctk.CTkFont(size=12), text_color="#a0a0a0").pack(pady=(0, 20))
    
    files_frame = ctk.CTkFrame(popup, fg_color="transparent")
    files_frame.pack(fill="x", padx=30)
    
    # Zirve Excel
    zirve_frame = ctk.CTkFrame(files_frame, fg_color="#363636", corner_radius=10)
    zirve_frame.grid(row=0, column=0, padx=(0, 10), sticky="nsew")
    files_frame.grid_columnconfigure(0, weight=1)
    
    ctk.CTkLabel(zirve_frame, text="Zirve Excel", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 5))
    zirve_status = ctk.CTkLabel(zirve_frame, text="📄 Dosya seçilmedi", font=ctk.CTkFont(size=11), text_color="#808080")
    zirve_status.pack(pady=(0, 10))
    
    def select_zirve():
        popup.withdraw()
        popup.update_idletasks()
        dosya = filedialog.askopenfilename(title="Zirve Excel Dosyası Seçin", filetypes=[("Excel Dosyaları", "*.xlsx *.xls")])
        popup.deiconify()
        popup.lift()
        popup.focus_force()
        if dosya:
            secilen_dosyalar["zirve"] = dosya
            dosya_adi = os.path.basename(dosya)
            if len(dosya_adi) > 20:
                dosya_adi = dosya_adi[:17] + "..."
            zirve_status.configure(text=f"✅ {dosya_adi}", text_color="#4CAF50")
            check_ready()
    
    ctk.CTkButton(zirve_frame, text="📂 Seç", width=100, height=32, command=select_zirve).pack(pady=(0, 15))
    
    # Portal Excel (Tek dosya)
    portal_frame = ctk.CTkFrame(files_frame, fg_color="#363636", corner_radius=10)
    portal_frame.grid(row=0, column=1, padx=(10, 0), sticky="nsew")
    files_frame.grid_columnconfigure(1, weight=1)
    
    ctk.CTkLabel(portal_frame, text="Gelen Fatura Excel", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 5))
    portal_status = ctk.CTkLabel(portal_frame, text="📄 Dosya seçilmedi", font=ctk.CTkFont(size=11), text_color="#808080")
    portal_status.pack(pady=(0, 10))
    
    def select_portal():
        popup.withdraw()
        popup.update_idletasks()
        dosya = filedialog.askopenfilename(title="Gelen Fatura Excel Dosyası Seçin", filetypes=[("Excel Dosyaları", "*.xlsx *.xls")])
        popup.deiconify()
        popup.lift()
        popup.focus_force()
        if dosya:
            secilen_dosyalar["portal"] = dosya
            dosya_adi = os.path.basename(dosya)
            if len(dosya_adi) > 18:
                dosya_adi = dosya_adi[:15] + "..."
            portal_status.configure(text=f"✅ {dosya_adi}", text_color="#4CAF50")
            check_ready()
    
    ctk.CTkButton(portal_frame, text="📂 Dosya Seç", width=110, height=32, command=select_portal).pack(pady=(0, 15))
    
    def run_kontrol():
        if secilen_dosyalar["zirve"] and secilen_dosyalar["portal"]:
            popup.destroy()
            hizlibilisim_alis_kontrol_callback(secilen_dosyalar["zirve"], secilen_dosyalar["portal"])
    
    kontrol_btn = ctk.CTkButton(popup, text="🔍 Kontrol Et", width=200, height=40, 
                                 command=run_kontrol, state="disabled",
                                 fg_color="#555555", hover_color="#555555")
    kontrol_btn.pack(pady=(25, 15))
    
    def check_ready():
        if secilen_dosyalar["zirve"] and secilen_dosyalar["portal"]:
            kontrol_btn.configure(state="normal", fg_color="#1f6aa5", hover_color="#144870")
    
    ctk.CTkButton(popup, text="İptal", width=100, height=30, 
                  fg_color="transparent", border_width=1, border_color="#555555",
                  hover_color="#404040", command=popup.destroy).pack()

def show_hizlibilisim_satis_popup():
    """Hızlıbilişim Satış kontrol popup'ı - çoklu dosya seçimi"""
    show_kontrol_popup("Hızlıbilişim Satış", hizlibilisim_satis_kontrol_callback)

def logo_hizli_otomatik_ayir_callback(zirve_path, portal_paths):
    """Portal dosyalarını otomatik tespit ederek Logo ve Hızlıbilişim olarak ayırır"""
    logo_portal_paths = []
    hizli_portal_paths = []
    
    # Her dosyayı oku ve tipini tespit et
    for path in portal_paths:
        try:
            df = pd.read_excel(path, nrows=0)  # Sadece sütun isimlerini oku
            sutunlar = df.columns.tolist()
            
            # Logo tespiti: "Toplam Tutar" ve "KDV Toplamı" varsa Logo
            if 'Toplam Tutar' in sutunlar and 'KDV Toplamı' in sutunlar:
                logo_portal_paths.append(path)
            # Hızlıbilişim tespiti: "FaturaNo" veya "PayableAmount" veya "OdenecekTutar" varsa Hızlı
            elif 'FaturaNo' in sutunlar or 'PayableAmount' in sutunlar or 'OdenecekTutar' in sutunlar:
                hizli_portal_paths.append(path)
            else:
                # Belirsiz, Logo olarak kabul et
                logo_portal_paths.append(path)
        except:
            # Hata durumunda Logo olarak kabul et
            logo_portal_paths.append(path)
    
    # Kombine karşılaştırma yap
    return logo_hizli_kombine_karsilastir(zirve_path, logo_portal_paths, hizli_portal_paths)

logo_hizli_kontrol_callback = _kontrol_sonuc_popup_goster(logo_hizli_otomatik_ayir_callback)

def show_logo_kontrol_popup():
    """Logo kontrol popup'ı - standart arayüz (arka planda Hızlıbilişim de kontrol edilir)"""
    show_kontrol_popup("Logo", logo_hizli_kontrol_callback)

def show_hizlibilisim_menu(button):
    """Hızlıbilişim için Alış/Satış popup menüsü göster"""
    global aktif_menu, popup_aktif
    popup_menu_kapat()
    
    x = button.winfo_rootx()
    y = button.winfo_rooty() + button.winfo_height() + 5
    
    aktif_menu = ctk.CTkToplevel(app)
    aktif_menu.geometry(f"160x90+{x}+{y}")
    aktif_menu.overrideredirect(True)
    aktif_menu.attributes("-topmost", True)
    aktif_menu.configure(fg_color="#3b3b3b")
    
    ctk.CTkButton(aktif_menu, text="📥 Alış Faturası", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), show_hizlibilisim_alis_popup()]).pack(pady=(8, 4), padx=5)
    ctk.CTkButton(aktif_menu, text="📤 Satış Faturası", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), show_hizlibilisim_satis_popup()]).pack(pady=(0, 8), padx=5)
    
    def activate_popup():
        global popup_aktif
        popup_aktif = True
    app.after(150, activate_popup)
    
    aktif_menu.bind("<Escape>", lambda e: popup_menu_kapat())

def show_uyumsoft_menu(button):
    """Uyumsoft için Zirve/Verim popup menüsü göster"""
    global aktif_menu, popup_aktif
    popup_menu_kapat()
    
    x = button.winfo_rootx()
    y = button.winfo_rooty() + button.winfo_height() + 5
    
    aktif_menu = ctk.CTkToplevel(app)
    aktif_menu.geometry(f"160x90+{x}+{y}")
    aktif_menu.overrideredirect(True)
    aktif_menu.attributes("-topmost", True)
    aktif_menu.configure(fg_color="#3b3b3b")
    
    ctk.CTkButton(aktif_menu, text="📊 Zirve", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), kontrol_uyumsoft_zirve()]).pack(pady=(8, 4), padx=5)
    ctk.CTkButton(aktif_menu, text="📊 Verim", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), kontrol_uyumsoft_verim()]).pack(pady=(0, 8), padx=5)
    
    def activate_popup():
        global popup_aktif
        popup_aktif = True
    app.after(150, activate_popup)
    
    aktif_menu.bind("<Escape>", lambda e: popup_menu_kapat())

# Satır 1: Logo ve Uyumsoft
kontrol_row1 = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_row1.pack(pady=8, expand=False)

logo_kontrol_col = ctk.CTkFrame(kontrol_row1, fg_color="transparent")
logo_kontrol_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(logo_kontrol_col, text="Logo", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(logo_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=show_logo_kontrol_popup).pack()

uyumsoft_kontrol_col = ctk.CTkFrame(kontrol_row1, fg_color="transparent")
uyumsoft_kontrol_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(uyumsoft_kontrol_col, text="Uyumsoft", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
uyumsoft_kontrol_btn = ctk.CTkButton(uyumsoft_kontrol_col, text="🔍 Kontrol", width=160, height=40)
uyumsoft_kontrol_btn.pack()
uyumsoft_kontrol_btn.configure(command=lambda: show_uyumsoft_menu(uyumsoft_kontrol_btn))

# Satır 2: Hızlıbilişim ve İzibiz
kontrol_row2 = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_row2.pack(pady=8, expand=False)

hizli_kontrol_col = ctk.CTkFrame(kontrol_row2, fg_color="transparent")
hizli_kontrol_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(hizli_kontrol_col, text="Hızlıbilişim", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
hizli_kontrol_btn = ctk.CTkButton(hizli_kontrol_col, text="🔍 Kontrol", width=160, height=40)
hizli_kontrol_btn.pack()
hizli_kontrol_btn.configure(command=lambda: show_hizlibilisim_menu(hizli_kontrol_btn))

izibiz_kontrol_col = ctk.CTkFrame(kontrol_row2, fg_color="transparent")
izibiz_kontrol_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(izibiz_kontrol_col, text="İzibiz", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(izibiz_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_izibiz).pack()

# Satır 3: Vega ve Hepsiburada/Trendyol
kontrol_row3 = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_row3.pack(pady=8, expand=False)

vega_kontrol_col = ctk.CTkFrame(kontrol_row3, fg_color="transparent")
vega_kontrol_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(vega_kontrol_col, text="Vega", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(vega_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_vega).pack()

hb_ve_trn_kontrol_col = ctk.CTkFrame(kontrol_row3, fg_color="transparent")
hb_ve_trn_kontrol_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(hb_ve_trn_kontrol_col, text="Hepsiburada/Trendyol", font=ctk.CTkFont(size=16)).pack(pady=(0, 10))
ctk.CTkButton(hb_ve_trn_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_hb_ve_trn).pack()

# Sonuç mesajı
kontrol_label_sonuc = ctk.CTkLabel(kontrol_inner, text="", font=ctk.CTkFont(size=16))
kontrol_label_sonuc.pack(pady=(20, 0))

# ================= BAŞLANGIÇ =================

# Ana pencereye tıklanınca popup menüyü kapat
app.bind_all("<Button-1>", on_app_click)

# Karşılama sayfasını göster
show_welcome()

app.mainloop()