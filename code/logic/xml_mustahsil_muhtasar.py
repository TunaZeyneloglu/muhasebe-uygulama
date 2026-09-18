import os
import glob
import re
import xml.etree.ElementTree as ET
from datetime import datetime
from tkinter import filedialog

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText

from constants import MASAUSTU, XML_NS
import state
import theme

# ----------------- Müstahsil Muhtasar (Zirve) -----------------

BASLIKLAR = [
    "Soyadı (F1)",
    "Adı (F1)",
    "Vergi Kimlik No (F1)",
    "T.C. Kimlik No",
    "Tür (F1)",
    "Gayrisafi Tutar",
    "Kesinti Tutar",
    "Belge Türü",
    "Belge Tarihi",
    "Belge Seri Sıra No",
    "Adresi",
    "Kdv Tevkifat Kodu (1, 2, 3)",
    "Adres No",
    "Ülke Kodu",
]

TUR_KODU = "093"
BELGE_TURU = "M. Makbuz"
TUTAR_FORMAT = '#,##0.00_);(#,##0.00)'
TARIH_FORMAT = 'dd\\.mm\\.yyyy'
METIN_FORMAT = '@'


def _to_float(deger):
    """Metni güvenli şekilde float'a çevir (virgüllü ondalık da desteklenir)."""
    if deger is None:
        return 0.0
    try:
        return float(str(deger).strip().replace(",", "."))
    except ValueError:
        return 0.0


def _bosluk_temizle(metin):
    """Fazla boşlukları teke indir ve baştaki/sondaki boşlukları at."""
    return re.sub(r"\s+", " ", (metin or "")).strip()


def _adres_olustur(party, ns):
    """Müşteri adresini oluştur: 'CitySubdivisionName CityName',
    ikisinden biri yoksa tamamen StreetName kullanılır."""
    adres_node = party.find("cac:PostalAddress", namespaces=ns)
    if adres_node is None:
        return ""
    ilce = _bosluk_temizle(adres_node.findtext("cbc:CitySubdivisionName", namespaces=ns))
    il = _bosluk_temizle(adres_node.findtext("cbc:CityName", namespaces=ns))
    if ilce and il:
        return _bosluk_temizle(f"{ilce} {il}")
    return _bosluk_temizle(adres_node.findtext("cbc:StreetName", namespaces=ns))


def _klasoru_isle(klasor):
    """GUI'den bağımsız saf işleme fonksiyonu.

    Klasördeki tüm XML dosyalarını (alt klasörler dahil) okur, müstahsil
    makbuzlarını doğrular ve Excel'e yazılacak satırları döndürür.
    """
    ns = XML_NS
    satirlar = []
    uyarilar = []
    atlananlar = []
    basarisiz_dosyalar = []
    gorulen_idler = {}
    dosya_sayisi = 0

    for xml_dosya in sorted(glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True)):
        dosya_adi = os.path.basename(xml_dosya)
        dosya_sayisi += 1
        try:
            root = ET.parse(xml_dosya).getroot()

            # Belge tipi kontrolü - müstahsil makbuzu olmayanlar atlanır
            belge_tipi = (root.findtext("cbc:CreditNoteTypeCode", namespaces=ns) or "").strip()
            if belge_tipi != "MUSTAHSILMAKBUZ":
                atlananlar.append(
                    f"{dosya_adi}: CreditNoteTypeCode '{belge_tipi or 'yok'}' (MUSTAHSILMAKBUZ değil)"
                )
                continue

            belge_no = (root.findtext("cbc:ID", namespaces=ns) or "").strip()

            # Aynı belge numarası birden fazla dosyada varsa ilki kalır
            if belge_no in gorulen_idler:
                atlananlar.append(
                    f"{dosya_adi}: {belge_no} numaralı belge zaten işlendi ({gorulen_idler[belge_no]}), tekrar atlandı"
                )
                uyarilar.append(f"{dosya_adi}: Mükerrer belge no {belge_no} - ilk dosya korundu")
                continue
            gorulen_idler[belge_no] = dosya_adi

            # Tarih
            tarih_raw = (root.findtext("cbc:IssueDate", namespaces=ns) or "").strip()
            try:
                belge_tarihi = datetime.strptime(tarih_raw, "%Y-%m-%d").date()
            except ValueError:
                belge_tarihi = None
                uyarilar.append(f"{dosya_adi}: Belge tarihi okunamadı ('{tarih_raw}')")

            # Müşteri (müstahsil) bilgileri
            musteri = root.find("cac:AccountingCustomerParty", namespaces=ns)
            if musteri is None:
                uyarilar.append(f"{dosya_adi}: cac:AccountingCustomerParty bulunamadı")
                soyadi = adi = vkn = tckn = ""
                adres = ""
            else:
                kisi = musteri.find(".//cac:Person", namespaces=ns)
                if kisi is None:
                    uyarilar.append(f"{dosya_adi}: cac:Person bulunamadı (Adı/Soyadı boş)")
                    soyadi = adi = ""
                else:
                    soyadi = (kisi.findtext("cbc:FamilyName", namespaces=ns) or "").strip()
                    adi = (kisi.findtext("cbc:FirstName", namespaces=ns) or "").strip()

                vkn = (musteri.findtext(
                    ".//cac:PartyIdentification/cbc:ID[@schemeID='VKN']", namespaces=ns) or "").strip()
                tckn = (musteri.findtext(
                    ".//cac:PartyIdentification/cbc:ID[@schemeID='TCKN']", namespaces=ns) or "").strip()
                if not tckn:
                    uyarilar.append(f"{dosya_adi}: TCKN bulunamadı")

                parti = musteri.find("cac:Party", namespaces=ns)
                adres = _adres_olustur(parti, ns) if parti is not None else ""

            if not adres:
                uyarilar.append(f"{dosya_adi}: Adres bilgisi boş")

            # Tutarlar
            gayrisafi = _to_float(
                root.findtext("cac:LegalMonetaryTotal/cbc:TaxExclusiveAmount", namespaces=ns))
            kesinti = _to_float(root.findtext("cac:TaxTotal/cbc:TaxAmount", namespaces=ns))

            # Stopaj oranı kontrolü (belge seviyesi TaxTotal)
            oran_raw = root.findtext(
                "cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
            if abs(_to_float(oran_raw) - 2.00) > 0.001:
                uyarilar.append(
                    f"{dosya_adi} ({belge_no}): Stopaj oranı %2 değil ('{(oran_raw or 'yok')}')")

            # Kesinti tutarı = gayrisafi * %2 kontrolü
            if abs(kesinti - gayrisafi * 0.02) > 0.01:
                uyarilar.append(
                    f"{dosya_adi} ({belge_no}): Kesinti tutarı uyuşmuyor "
                    f"(beklenen {gayrisafi * 0.02:.2f}, belgede {kesinti:.2f})")

            # Satır toplamı = belge LineExtensionAmount kontrolü
            satir_toplami = sum(
                _to_float(satir.findtext("cbc:LineExtensionAmount", namespaces=ns))
                for satir in root.findall("cac:CreditNoteLine", namespaces=ns)
            )
            belge_satir_toplami = _to_float(
                root.findtext("cac:LegalMonetaryTotal/cbc:LineExtensionAmount", namespaces=ns))
            if abs(satir_toplami - belge_satir_toplami) > 0.01:
                uyarilar.append(
                    f"{dosya_adi} ({belge_no}): Satır toplamı belge toplamıyla uyuşmuyor "
                    f"(satırlar {satir_toplami:.2f}, belge {belge_satir_toplami:.2f})")

            satirlar.append({
                "soyadi": soyadi,
                "adi": adi,
                "vkn": vkn,
                "tckn": tckn,
                "gayrisafi": gayrisafi,
                "kesinti": kesinti,
                "belge_tarihi": belge_tarihi,
                "belge_no": belge_no,
                "adres": adres,
                "dosya": dosya_adi,
            })

        except Exception as e:
            basarisiz_dosyalar.append(dosya_adi)
            uyarilar.append(f"{dosya_adi}: Dosya okunamadı - {e}")

    # Sıralama: Belge tarihi, sonra belge no
    satirlar.sort(key=lambda s: (
        s["belge_tarihi"] or datetime.min.date(),
        s["belge_no"],
    ))

    return {
        "satirlar": satirlar,
        "uyarilar": uyarilar,
        "atlananlar": atlananlar,
        "dosya_sayisi": dosya_sayisi,
        "basarisiz_dosyalar": basarisiz_dosyalar,
    }


def _metin_yaz(hucre, deger):
    """Hücreye metin yaz. openpyxl düz '' değerini kaydederken boş (None) hücreye
    çevirdiği için, boş metinler CellRichText('') ile yazılır; böylece dosya geri
    okunduğunda değer None değil '' olarak gelir."""
    metin = "" if deger is None else str(deger)
    hucre.value = CellRichText("") if metin == "" else metin


def _excel_yaz(satirlar, dosya_yolu):
    """Satırları 'Zirve' sayfasına, sütun tiplerini koruyarak yazar."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Zirve"

    for kolon, baslik in enumerate(BASLIKLAR, start=1):
        ws.cell(row=1, column=kolon).value = baslik

    for indis, satir in enumerate(satirlar):
        r = indis + 2

        ws.cell(row=r, column=1).value = satir["soyadi"]
        ws.cell(row=r, column=2).value = satir["adi"]
        _metin_yaz(ws.cell(row=r, column=3), satir["vkn"])

        tckn_hucre = ws.cell(row=r, column=4)
        _metin_yaz(tckn_hucre, satir["tckn"])
        tckn_hucre.number_format = METIN_FORMAT

        tur_hucre = ws.cell(row=r, column=5)
        tur_hucre.value = str(TUR_KODU)
        tur_hucre.number_format = METIN_FORMAT

        gayrisafi_hucre = ws.cell(row=r, column=6)
        gayrisafi_hucre.value = float(satir["gayrisafi"])
        gayrisafi_hucre.number_format = TUTAR_FORMAT

        kesinti_hucre = ws.cell(row=r, column=7)
        kesinti_hucre.value = float(satir["kesinti"])
        kesinti_hucre.number_format = TUTAR_FORMAT

        ws.cell(row=r, column=8).value = BELGE_TURU

        tarih_hucre = ws.cell(row=r, column=9)
        tarih_hucre.value = satir["belge_tarihi"]
        tarih_hucre.number_format = TARIH_FORMAT

        belge_no_hucre = ws.cell(row=r, column=10)
        _metin_yaz(belge_no_hucre, satir["belge_no"])
        belge_no_hucre.number_format = METIN_FORMAT

        _metin_yaz(ws.cell(row=r, column=11), satir["adres"])
        _metin_yaz(ws.cell(row=r, column=12), '')
        # M sütunu (Adres No) bilinçli olarak yazılmaz - boş (None) kalır
        _metin_yaz(ws.cell(row=r, column=14), '')

    wb.save(dosya_yolu)


def xml_mustahsil_muhtasar_oku_ve_yaz(klasor, label_sonuc, buton_ac):
    sonuc = _klasoru_isle(klasor)
    satirlar = sonuc["satirlar"]
    uyarilar = sonuc["uyarilar"]
    atlananlar = sonuc["atlananlar"]
    basarisiz_dosyalar = sonuc["basarisiz_dosyalar"]

    # Konsol özeti
    print(
        f"{sonuc['dosya_sayisi']} dosya okundu, {len(satirlar)} satır yazıldı, "
        f"{len(atlananlar)} atlandı."
    )
    if atlananlar:
        print("Atlananlar:")
        for atlanan in atlananlar:
            print(f"  - {atlanan}")
    if uyarilar:
        print(f"Uyarılar ({len(uyarilar)}):")
        for uyari in uyarilar:
            print(f"  - {uyari}")

    if satirlar:
        dosya_yolu = os.path.join(MASAUSTU, "muhtasar_zirve.xlsx")
        _excel_yaz(satirlar, dosya_yolu)
        if basarisiz_dosyalar:
            label_sonuc.configure(
                text=f"muhtasar_zirve.xlsx kaydedildi, ancak {len(basarisiz_dosyalar)} dosya işlenemedi: {', '.join(basarisiz_dosyalar[:3])}{' ...' if len(basarisiz_dosyalar) > 3 else ''}",
                text_color=theme.WARNING
            )
        elif atlananlar or uyarilar:
            label_sonuc.configure(
                text=f"muhtasar_zirve.xlsx kaydedildi ({len(satirlar)} satır), {len(atlananlar)} belge atlandı, {len(uyarilar)} uyarı (detaylar konsolda).",
                text_color=theme.WARNING
            )
        else:
            label_sonuc.configure(
                text=f"muhtasar_zirve.xlsx masaüstüne kaydedildi ({len(satirlar)} satır).",
                text_color=theme.SUCCESS
            )
        buton_ac.configure(state="normal")
        state.SON_DOSYA_YOLU = dosya_yolu
    elif basarisiz_dosyalar:
        label_sonuc.configure(
            text=f"Geçerli müstahsil makbuzu bulunamadı. ({len(basarisiz_dosyalar)} dosya hata verdi)",
            text_color=theme.ERROR
        )
    else:
        label_sonuc.configure(
            text="Geçerli müstahsil makbuzu bulunamadı.",
            text_color=theme.ERROR
        )


def klasor_sec_ve_isle_mustahsil_muhtasar(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Müstahsil makbuz XML klasörünü seçin")
    if secilen_klasor:
        xml_mustahsil_muhtasar_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac)
