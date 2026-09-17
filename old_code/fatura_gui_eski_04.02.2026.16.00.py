import os
import glob
import xml.etree.ElementTree as ET
import pandas as pd
from datetime import datetime
import subprocess
import sys
import unicodedata
from pathlib import Path
import customtkinter as ctk
from tkinter import filedialog

# Tema ayarı
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

MASAUSTU = str(Path.home() / "OneDrive" / "Desktop")
if not os.path.exists(MASAUSTU):
    MASAUSTU = str(Path.home() / "Desktop")

SON_DOSYA_YOLU = None
TARGET_VKN_FOR_ZERO_VAT = "7740042326"

# ----------------- Ortak Fonksiyonlar -----------------

def cari_adi_al(party, ns):
    cariadi = party.findtext("cac:PartyName/cbc:Name", namespaces=ns)
    if not cariadi:
        firstname = party.findtext("cac:Person/cbc:FirstName", namespaces=ns)
        lastname = party.findtext("cac:Person/cbc:FamilyName", namespaces=ns)
        if firstname and lastname:
            cariadi = f"{firstname} {lastname}"
    if not cariadi:
        cariadi = party.findtext("cac:PartyLegalEntity/cbc:RegistrationName", namespaces=ns)
    return cariadi if cariadi else "Bilinmeyen Cari"

def kaydet_excel(veriler, dosya_yolu):
    df = pd.DataFrame(veriler)[[
        "SIRANO", "TARIH", "EVRAKNO", "EVRAKTURU", "CARIKODU", "CARIADI",
        "STOKKODU", "STOKADI", "BIRIM", "MIKTAR", "DOVIZC",
        "BIRIMFIYAT", "INDTL", "KDVY", "TUTARTL", "FATURACK"
    ]]
    df["EVRAKNO"] = df["EVRAKNO"].astype(str)
    with pd.ExcelWriter(dosya_yolu, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Sayfa1")

def dosyayi_ac():
    global SON_DOSYA_YOLU
    if SON_DOSYA_YOLU:
        try:
            if sys.platform == "win32":
                os.startfile(SON_DOSYA_YOLU)
            elif sys.platform == "darwin":
                subprocess.call(["open", SON_DOSYA_YOLU])
            else:
                subprocess.call(["xdg-open", SON_DOSYA_YOLU])
        except:
            pass

# ----------------- Uyumsoft Kontrol Fonksiyonu -----------------

def uyumsoft_karsilastir(zirve_path, portal_paths):
    """
    Zirve ve Uyumsoft Portal Excel dosyalarını karşılaştırır.
    
    Args:
        zirve_path: Tek Zirve Excel dosyası yolu
        portal_paths: Portal Excel dosyası yolu (string) veya dosya yolları listesi
    
    Returns:
        tuple: (sonuc_dosyasi_yolu, mesaj) veya (None, hata_mesaji)
    """
    global SON_DOSYA_YOLU
    
    try:
        # Excel dosyalarını oku
        df_zirve = pd.read_excel(zirve_path)
        
        # Portal dosyaları - liste veya tek dosya olabilir
        if isinstance(portal_paths, list):
            # Çoklu dosya - hepsini oku ve birleştir
            portal_dfs = []
            for path in portal_paths:
                df_temp = pd.read_excel(path)
                portal_dfs.append(df_temp)
            df_portal = pd.concat(portal_dfs, ignore_index=True)
        else:
            # Tek dosya
            df_portal = pd.read_excel(portal_paths)
        
        # Fatura No kolonlarını temizle ve normalize et
        def temizle_fatura_no(fatura_no):
            if pd.isna(fatura_no):
                return ""
            # Apostrof ve boşlukları kaldır
            return str(fatura_no).strip().replace("'", "").upper()
        
        df_zirve['Fatura_No_Temiz'] = df_zirve['Fatura No'].apply(temizle_fatura_no)
        df_portal['Fatura_No_Temiz'] = df_portal['Fatura No'].apply(temizle_fatura_no)
        
        # Tutarları float'a çevir
        df_zirve['Tutar_TL'] = pd.to_numeric(df_zirve['Tutar TL'], errors='coerce').fillna(0)
        df_zirve['KDV_TL'] = pd.to_numeric(df_zirve['KDV TL'], errors='coerce').fillna(0)
        
        df_portal['Odenecek_Tutar'] = pd.to_numeric(df_portal['Ödenecek Tutar'], errors='coerce').fillna(0)
        df_portal['Toplam_KDV'] = pd.to_numeric(df_portal['Toplam KDV'], errors='coerce').fillna(0)
        
        # Set'ler oluştur
        zirve_faturalar = set(df_zirve['Fatura_No_Temiz'])
        portal_faturalar = set(df_portal['Fatura_No_Temiz'])
        
        # 1. TUTAR FARKLARI (Her iki tarafta da var)
        tutar_farklari = []
        ortak_faturalar = zirve_faturalar.intersection(portal_faturalar)
        
        for fatura_no in ortak_faturalar:
            if not fatura_no:  # Boş fatura no'ları atla
                continue
                
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            tutar_farki = abs(zirve_row['Tutar_TL'] - portal_row['Odenecek_Tutar'])
            kdv_farki = abs(zirve_row['KDV_TL'] - portal_row['Toplam_KDV'])
            
            # 1 TL'den fazla fark varsa kaydet
            if tutar_farki > 1 or kdv_farki > 1:
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Zirve Cari': zirve_row.get('Cari Unvanı', ''),
                    'Portal Cari': portal_row.get('Gönderici', ''),
                    'Zirve Tutar': round(zirve_row['Tutar_TL'], 2),
                    'Portal Tutar': round(portal_row['Odenecek_Tutar'], 2),
                    'Tutar Farkı': round(tutar_farki, 2),
                    'Zirve KDV': round(zirve_row['KDV_TL'], 2),
                    'Portal KDV': round(portal_row['Toplam_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Zirve Tarih': zirve_row.get('Tarih', ''),
                    'Portal Tarih': portal_row.get('Fatura Tarihi', '')
                })
        
        # 2. EKSİK GİRİŞLER (Portal'de var, Zirve'de yok)
        eksik_girisler = []
        eksik_faturalar = portal_faturalar - zirve_faturalar
        
        for fatura_no in eksik_faturalar:
            if not fatura_no:
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'Gönderici VKN/TCKN': portal_row.get('Gönderici VKN/TCKN', ''),
                'Gönderici': portal_row.get('Gönderici', ''),
                'Ödenecek Tutar': round(portal_row['Odenecek_Tutar'], 2),
                'Toplam KDV': round(portal_row['Toplam_KDV'], 2),
                'Fatura Tarihi': portal_row.get('Fatura Tarihi', ''),
                'Oluşturulma Tarihi': portal_row.get('Oluşturulma Tarihi', '')
            })
        
        # 3. FAZLA GİRİŞLER (Zirve'de var, Portal'de yok)
        fazla_girisler = []
        fazla_faturalar = zirve_faturalar - portal_faturalar
        
        for fatura_no in fazla_faturalar:
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': zirve_row.get('Cari Unvanı', ''),
                'Tutar TL': round(zirve_row['Tutar_TL'], 2),
                'KDV TL': round(zirve_row['KDV_TL'], 2),
                'Tarih': zirve_row.get('Tarih', ''),
                'E.Kod': zirve_row.get('E.Kod', ''),
                'İşlem Türü': zirve_row.get('İşlem Türü', '')
            })
        
        # Excel dosyası oluştur
        sonuc_dosyasi = os.path.join(MASAUSTU, "Uyumsoft_Kontrol_Sonuc.xlsx")
        
        with pd.ExcelWriter(sonuc_dosyasi, engine='openpyxl') as writer:
            # Sayfa 1: Tutar Farkları
            if tutar_farklari:
                df_tutar = pd.DataFrame(tutar_farklari)
                # Fatura No'ya göre sırala
                df_tutar = df_tutar.sort_values('Fatura No')
                df_tutar.to_excel(writer, sheet_name='Tutar Farkları', index=False)
                
                # Formatlama - Tutar Farkı ve KDV Farkı sütunlarını kalın ve büyük yap
                worksheet = writer.sheets['Tutar Farkları']
                from openpyxl.styles import Font
                
                # Sütun indekslerini bul
                headers = list(df_tutar.columns)
                tutar_farki_col = headers.index('Tutar Farkı') + 1  # Excel 1-indexed
                kdv_farki_col = headers.index('KDV Farkı') + 1
                
                # Her satır için formatlama (başlık hariç)
                for row_idx in range(2, len(df_tutar) + 2):
                    worksheet.cell(row=row_idx, column=tutar_farki_col).font = Font(bold=True, size=13)
                    worksheet.cell(row=row_idx, column=kdv_farki_col).font = Font(bold=True, size=13)
            else:
                pd.DataFrame({'Mesaj': ['Tutar farkı bulunamadı.']}).to_excel(
                    writer, sheet_name='Tutar Farkları', index=False)
            
            # Sayfa 2: Eksik Girişler
            if eksik_girisler:
                df_eksik = pd.DataFrame(eksik_girisler)
                # Fatura No'ya göre sırala
                df_eksik = df_eksik.sort_values('Fatura No')
                df_eksik.to_excel(writer, sheet_name='Eksik Girişler', index=False)
            else:
                pd.DataFrame({'Mesaj': ['Eksik giriş bulunamadı.']}).to_excel(
                    writer, sheet_name='Eksik Girişler', index=False)
            
            # Sayfa 3: Fazla Girişler
            if fazla_girisler:
                df_fazla = pd.DataFrame(fazla_girisler)
                # Fatura No'ya göre sırala
                df_fazla = df_fazla.sort_values('Fatura No')
                df_fazla.to_excel(writer, sheet_name='Fazla Girişler', index=False)
            else:
                pd.DataFrame({'Mesaj': ['Fazla giriş bulunamadı.']}).to_excel(
                    writer, sheet_name='Fazla Girişler', index=False)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        
        # Özet mesajı oluştur
        portal_dosya_sayisi = len(portal_paths) if isinstance(portal_paths, list) else 1
        mesaj = f"✅ Kontrol tamamlandı!\n\n"
        mesaj += f"📁 İşlenen Portal Dosyası: {portal_dosya_sayisi}\n"
        mesaj += f"📄 Toplam Portal Fatura: {len(df_portal)}\n"
        mesaj += f"📄 Toplam Zirve Fatura: {len(df_zirve)}\n\n"
        mesaj += f"📊 Tutar Farkları: {len(tutar_farklari)}\n"
        mesaj += f"⚠️ Eksik Girişler: {len(eksik_girisler)}\n"
        mesaj += f"❌ Fazla Girişler: {len(fazla_girisler)}\n\n"
        mesaj += f"Sonuç dosyası masaüstüne kaydedildi."
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        return (None, f"❌ Hata oluştu:\n{str(e)}")

# ----------------- Vega Kontrol Fonksiyonu -----------------

def vega_karsilastir(zirve_path, portal_paths):
    """
    Zirve ve Vega Portal Excel dosyalarını karşılaştırır.
    
    Args:
        zirve_path: Tek Zirve Excel dosyası yolu
        portal_paths: Portal Excel dosyası yolu (string) veya dosya yolları listesi
    
    Returns:
        tuple: (sonuc_dosyasi_yolu, mesaj) veya (None, hata_mesaji)
    """
    global SON_DOSYA_YOLU
    
    try:
        # Excel dosyalarını oku
        df_zirve = pd.read_excel(zirve_path)
        
        # Portal dosyaları - liste veya tek dosya olabilir
        if isinstance(portal_paths, list):
            # Çoklu dosya - hepsini oku ve birleştir
            portal_dfs = []
            for path in portal_paths:
                df_temp = pd.read_excel(path)
                portal_dfs.append(df_temp)
            df_portal = pd.concat(portal_dfs, ignore_index=True)
        else:
            # Tek dosya
            df_portal = pd.read_excel(portal_paths)
        
        # Fatura No kolonlarını temizle ve normalize et
        def temizle_fatura_no(fatura_no):
            if pd.isna(fatura_no):
                return ""
            # Apostrof ve boşlukları kaldır
            return str(fatura_no).strip().replace("'", "").upper()
        
        # Tutarları sayıya çevir (virgül ve noktayı temizle)
        def parse_tutar(value):
            if pd.isna(value):
                return 0
            if isinstance(value, (int, float)):
                return float(value)
            # String ise virgül ve noktayı temizle
            return float(str(value).replace('.', '').replace(',', '.'))
        
        df_zirve['Fatura_No_Temiz'] = df_zirve['Fatura No'].apply(temizle_fatura_no)
        df_portal['Fatura_No_Temiz'] = df_portal['Fatura No'].apply(temizle_fatura_no)
        
        # Zirve tutarları - Tutar TL zaten KDV dahil toplam tutar
        df_zirve['Toplam_TL'] = df_zirve['Tutar TL'].apply(parse_tutar)  # Zaten KDV dahil
        df_zirve['KDV_TL'] = df_zirve['KDV TL'].apply(parse_tutar)
        
        # Portal tutarları - farklı sütun isimleri olabilir, her satır için doğru sütunu kullan
        def get_portal_toplam(row):
            """Ödenecek veya Toplam sütununu al"""
            if 'Ödenecek' in df_portal.columns and pd.notna(row.get('Ödenecek')):
                return parse_tutar(row['Ödenecek'])
            elif 'Toplam' in df_portal.columns and pd.notna(row.get('Toplam')):
                return parse_tutar(row['Toplam'])
            return 0
        
        df_portal['Portal_KDV'] = df_portal['Vergi'].apply(parse_tutar)
        df_portal['Portal_Toplam'] = df_portal.apply(get_portal_toplam, axis=1)
        
        # Set'ler oluştur
        zirve_faturalar = set(df_zirve['Fatura_No_Temiz']) - {''}
        portal_faturalar = set(df_portal['Fatura_No_Temiz']) - {''}
        
        # 1. TUTAR FARKLARI (Her iki tarafta da var) - Uyumsoft formatı ile aynı sıralama
        tutar_farklari = []
        ortak_faturalar = zirve_faturalar.intersection(portal_faturalar)
        
        for fatura_no in ortak_faturalar:
            if not fatura_no:  # Boş fatura no'ları atla
                continue
                
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            
            # Toplam (KDV dahil) ve KDV tutarları karşılaştır
            toplam_farki = abs(zirve_row['Toplam_TL'] - portal_row['Portal_Toplam'])
            kdv_farki = abs(zirve_row['KDV_TL'] - portal_row['Portal_KDV'])
            
            # 1 TL'den fazla fark varsa kaydet (Uyumsoft formatı ile aynı sıralama)
            if toplam_farki > 1 or kdv_farki > 1:
                tutar_farklari.append({
                    'Fatura No': fatura_no,
                    'Zirve Cari': zirve_row.get('Cari Unvanı', ''),
                    'Portal Cari': portal_row.get('Unvan', ''),
                    'Zirve Tutar': round(zirve_row['Toplam_TL'], 2),
                    'Portal Tutar': round(portal_row['Portal_Toplam'], 2),
                    'Tutar Farkı': round(toplam_farki, 2),
                    'Zirve KDV': round(zirve_row['KDV_TL'], 2),
                    'Portal KDV': round(portal_row['Portal_KDV'], 2),
                    'KDV Farkı': round(kdv_farki, 2),
                    'Zirve Tarih': zirve_row.get('Tarih', ''),
                    'Portal Tarih': portal_row.get('Tarih', '')
                })
        
        # 2. EKSİK GİRİŞLER (Portal'de var, Zirve'de yok) - Uyumsoft formatı ile aynı
        eksik_girisler = []
        eksik_faturalar = portal_faturalar - zirve_faturalar
        
        for fatura_no in eksik_faturalar:
            if not fatura_no:
                continue
            portal_row = df_portal[df_portal['Fatura_No_Temiz'] == fatura_no].iloc[0]
            eksik_girisler.append({
                'Fatura No': fatura_no,
                'VKN/TCKN': portal_row.get('VKN/TCKN', ''),
                'Unvan': portal_row.get('Unvan', ''),
                'Ödenecek Tutar': round(portal_row['Portal_Toplam'], 2),
                'Toplam KDV': round(portal_row['Portal_KDV'], 2),
                'Fatura Tarihi': portal_row.get('Tarih', ''),
                'Fatura Türü': portal_row.get('Türü', portal_row.get('Fatura Türü', ''))
            })
        
        # 3. FAZLA GİRİŞLER (Zirve'de var, Portal'de yok) - Uyumsoft formatı ile aynı
        fazla_girisler = []
        fazla_faturalar = zirve_faturalar - portal_faturalar
        
        for fatura_no in fazla_faturalar:
            if not fatura_no:
                continue
            zirve_row = df_zirve[df_zirve['Fatura_No_Temiz'] == fatura_no].iloc[0]
            fazla_girisler.append({
                'Fatura No': fatura_no,
                'Cari Unvanı': zirve_row.get('Cari Unvanı', ''),
                'Tutar TL': round(zirve_row['Toplam_TL'], 2),
                'KDV TL': round(zirve_row['KDV_TL'], 2),
                'Tarih': zirve_row.get('Tarih', ''),
                'E.Kod': zirve_row.get('E.Kod', ''),
                'İşlem Türü': zirve_row.get('İşlem Türü', '')
            })
        
        # Excel dosyası oluştur
        sonuc_dosyasi = os.path.join(MASAUSTU, "Vega_Kontrol_Sonuc.xlsx")
        
        with pd.ExcelWriter(sonuc_dosyasi, engine='openpyxl') as writer:
            # Sayfa 1: Tutar Farkları - Uyumsoft formatı ile aynı
            if tutar_farklari:
                df_tutar = pd.DataFrame(tutar_farklari)
                # Fatura No'ya göre sırala
                df_tutar = df_tutar.sort_values('Fatura No')
                df_tutar.to_excel(writer, sheet_name='Tutar Farkları', index=False)
                
                # Formatlama - Tutar Farkı ve KDV Farkı sütunlarını kalın ve büyük yap
                worksheet = writer.sheets['Tutar Farkları']
                from openpyxl.styles import Font
                
                # Sütun indekslerini bul
                headers = list(df_tutar.columns)
                tutar_farki_col = headers.index('Tutar Farkı') + 1  # Excel 1-indexed
                kdv_farki_col = headers.index('KDV Farkı') + 1
                
                # Her satır için formatlama (başlık hariç)
                for row_idx in range(2, len(df_tutar) + 2):
                    worksheet.cell(row=row_idx, column=tutar_farki_col).font = Font(bold=True, size=13)
                    worksheet.cell(row=row_idx, column=kdv_farki_col).font = Font(bold=True, size=13)
            else:
                pd.DataFrame({'Mesaj': ['Tutar farkı bulunamadı.']}).to_excel(
                    writer, sheet_name='Tutar Farkları', index=False)
            
            # Sayfa 2: Eksik Girişler
            if eksik_girisler:
                df_eksik = pd.DataFrame(eksik_girisler)
                df_eksik = df_eksik.sort_values('Fatura No')
                df_eksik.to_excel(writer, sheet_name='Eksik Girişler', index=False)
            else:
                pd.DataFrame({'Mesaj': ['Eksik giriş bulunamadı.']}).to_excel(
                    writer, sheet_name='Eksik Girişler', index=False)
            
            # Sayfa 3: Fazla Girişler
            if fazla_girisler:
                df_fazla = pd.DataFrame(fazla_girisler)
                df_fazla = df_fazla.sort_values('Fatura No')
                df_fazla.to_excel(writer, sheet_name='Fazla Girişler', index=False)
            else:
                pd.DataFrame({'Mesaj': ['Fazla giriş bulunamadı.']}).to_excel(
                    writer, sheet_name='Fazla Girişler', index=False)
        
        SON_DOSYA_YOLU = sonuc_dosyasi
        
        # Özet mesajı oluştur
        portal_dosya_sayisi = len(portal_paths) if isinstance(portal_paths, list) else 1
        mesaj = f"✅ Kontrol tamamlandı!\n\n"
        mesaj += f"📁 İşlenen Portal Dosyası: {portal_dosya_sayisi}\n"
        mesaj += f"📄 Toplam Portal Fatura: {len(df_portal)}\n"
        mesaj += f"📄 Toplam Zirve Fatura: {len(df_zirve)}\n\n"
        mesaj += f"📊 Tutar Farkları: {len(tutar_farklari)}\n"
        mesaj += f"⚠️ Eksik Girişler: {len(eksik_girisler)}\n"
        mesaj += f"❌ Fazla Girişler: {len(fazla_girisler)}\n\n"
        mesaj += f"Sonuç dosyası masaüstüne kaydedildi."
        
        return (sonuc_dosyasi, mesaj)
        
    except Exception as e:
        import traceback
        return (None, f"❌ Hata oluştu:\n{str(e)}\n\n{traceback.format_exc()}")

# ----------------- Logo XML -----------------

def xml_logo_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    global SON_DOSYA_YOLU
    veriler = []
    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            ns = {
                "cbc": "urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2",
                "cac": "urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2",
            }
            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)
            carikodu = root.findtext(".//cac:AccountingCustomerParty//cbc:ID", namespaces=ns)
            party = root.find(".//cac:AccountingCustomerParty//cac:Party", namespaces=ns)
            cariadi = cari_adi_al(party, ns) if party is not None else "Bilinmeyen Cari"
            doviz = root.findtext(".//cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"

            for line in root.findall(".//cac:InvoiceLine", namespaces=ns):
                miktar = int(float(line.findtext("cbc:InvoicedQuantity", default="0", namespaces=ns)))
                birim = "ADET"
                fiyat = float(line.findtext("cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext("cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext("cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns)

                iskonto_element = line.find("cac:AllowanceCharge/cbc:Amount", namespaces=ns)
                iskonto = float(iskonto_element.text.replace(",", ".")) if iskonto_element is not None else 0.0

                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except:
                    kdv_orani = 0

                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Satış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": birim,
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto,
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")

    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, f"{fatura_tipi}.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text=f"✅ {fatura_tipi}.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_ve_isle_logo(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Logo XML klasörünü seçin")
    if secilen_klasor:
        xml_logo_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "logo_faturalar")

# ----------------- Uyumsoft XML -----------------

def xml_uyumsoft_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    global SON_DOSYA_YOLU
    veriler = []

    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()

            ns = {
                "cbc": "urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2",
                "cac": "urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2",
            }

            # Tarih ve Evrak No
            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)

            # VKN kontrolü - eğer Satıcı VKN'si hedef VKN ise KDV sıfırlanacak
            target_vkn = root.findtext(".//cac:AccountingSupplierParty//cbc:ID[@schemeID='VKN']", namespaces=ns)
            force_kdv_zero = (target_vkn == TARGET_VKN_FOR_ZERO_VAT)

            # Satıcı VKN ve Adı
            supplier_party = root.find(".//cac:AccountingSupplierParty/cac:Party", namespaces=ns)
            carikodu = supplier_party.findtext("cac:PartyIdentification/cbc:ID", namespaces=ns)
            cariadi = cari_adi_al(supplier_party, ns)

            # Yardımcı fonksiyon: güvenli float dönüşümü
            def _to_float(x):
                return float(str(x).replace(",", ".")) if x is not None else 0.0

            # KDV'ye göre gruplama
            kdv_dict = {}
            lines = root.findall(".//cac:InvoiceLine", namespaces=ns)

            for line in lines:
                # BaseAmount (önce AllowanceCharge.BaseAmount'u dene, yoksa LineExtensionAmount)
                base_amount = 0.0
                ac_nodes = line.findall("cac:AllowanceCharge", namespaces=ns)
                if ac_nodes:
                    for ac in ac_nodes:
                        ba = ac.findtext("cbc:BaseAmount", namespaces=ns)
                        if ba is not None:
                            base_amount = _to_float(ba)
                            break
                if base_amount == 0.0:
                    base_amount = _to_float(line.findtext("cbc:LineExtensionAmount", namespaces=ns))

                # İndirim toplamı: ChargeIndicator == "false" olanların Amount toplamı
                discount_total = 0.0
                for ac in ac_nodes:
                    ch = (ac.findtext("cbc:ChargeIndicator", namespaces=ns) or "").strip().lower()
                    if ch == "false":
                        discount_total += _to_float(ac.findtext("cbc:Amount", namespaces=ns))

                # Satır değeri hesapla (KDV hariç)
                line_value = base_amount - discount_total
                if line_value < 0:
                    line_value = 0.0

                # KDV oranı
                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except:
                    kdv_orani = 0

                # VKN'ye göre KDV sıfırlama
                if force_kdv_zero:
                    kdv_orani = 0

                # Grupla
                kdv_dict[kdv_orani] = kdv_dict.get(kdv_orani, 0.0) + line_value

            # Excel için veri satırlarını oluştur
            for kdv, toplam in kdv_dict.items():
                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Alış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": str(kdv),    # KDV oranı = StokKodu
                    "STOKADI": "",           # Her zaman boş
                    "BIRIM": "ADET",         # Sabit
                    "MIKTAR": 1,             # Sabit
                    "DOVIZC": "TL",          # Sabit
                    "BIRIMFIYAT": round(toplam, 2),
                    "INDTL": "",             # Boş
                    "KDVY": kdv,
                    "TUTARTL": round(toplam, 2),
                    "FATURACK": "",
                })

        except Exception as e:
            label_sonuc.configure(
                text=f"HATA: {os.path.basename(xml_dosya)} - {e}",
                text_color="red"
            )

    # Kaydetme işlemi
    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, f"{fatura_tipi}.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(
            text=f"✅ {fatura_tipi}.xlsx masaüstüne kaydedildi.",
            text_color="green"
        )
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(
            text="Geçerli XML dosyası bulunamadı.",
            text_color="red"
        )

def klasor_sec_ve_isle_uyumsoft(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Uyumsoft XML klasörünü seçin")
    if secilen_klasor:
        xml_uyumsoft_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "uyumsoft_faturalar")

# ----------------- Müstahsil XML -----------------

def temizle_kelime(kelime):
    kelime = kelime.strip().lower()
    kelime = kelime.replace("ı", "i").replace("ş", "s").replace("ç", "c").replace("ğ", "g").replace("ü", "u").replace("ö", "o")
    kelime = unicodedata.normalize("NFD", kelime).encode("ascii", "ignore").decode("utf-8")
    return kelime

def xml_mustahsil_oku_ve_yaz(klasor, label_sonuc, buton_ac, fatura_tipi):
    global SON_DOSYA_YOLU

    urun_kodlari = {
        "patates": "08120101",
        "kuru sogan": "08110101",
        "sarimsak": "08190101",
        "limon": "16030301"
    }

    ns = {
        "cbc": "urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2",
        "cac": "urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2",
    }

    veriler = []

    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()

            belge_no = root.findtext(".//cbc:ID", namespaces=ns)
            belge_tarih = root.findtext(".//cbc:IssueDate", namespaces=ns)

            # TC kimlik no
            tc_no = ""
            for elem in root.findall(".//cac:AccountingCustomerParty//cbc:ID", namespaces=ns):
                if elem.attrib.get("schemeID") == "TCKN":
                    tc_no = elem.text.strip()
                    break

            for line in root.findall(".//cac:CreditNoteLine", namespaces=ns):
                miktar = line.findtext(".//cbc:CreditedQuantity", namespaces=ns)
                birim_fiyat = line.findtext(".//cbc:PriceAmount", namespaces=ns)
                tutar = line.findtext(".//cbc:LineExtensionAmount", namespaces=ns)
                urun_adi_raw = line.findtext(".//cac:Item//cbc:Name", namespaces=ns)

                if not urun_adi_raw:
                    continue

                urun_adi = temizle_kelime(urun_adi_raw)
                urun_kodu = urun_kodlari.get(urun_adi, "")

                veriler.append({
                    "kimlik tipi": "k",
                    "TC KİMLİK NO": tc_no,
                    "MİKTAR": int(float(miktar)) if miktar else "",
                    "BİRİM FİYATI": birim_fiyat.replace(".", ",") if birim_fiyat else "",
                    "SATIS_TUTAR": tutar.replace(".", ",") if tutar else "",
                    "BELGE NO": belge_no,
                    "BELGE TARİHİ": f"{int(belge_tarih[8:10])}/{int(belge_tarih[5:7])}/{belge_tarih[:4]}" if belge_tarih else "",
                    "ÜRÜN KODU": urun_kodu,
                    "MENŞE": "001",
                    "MAHSÜL YILI": "2025",
                    "SATIŞ ŞEKLİ": "011",
                    "BAĞKUR MİKTAR": "",
                    "YRD_MIKTAR": "",
                    "YRD_OLCU_BIRIM": ""
                })

        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")

    if veriler:
        df = pd.DataFrame(veriler)
        dosya_yolu = os.path.join(MASAUSTU, f"{fatura_tipi}.xlsx")
        df.to_excel(dosya_yolu, index=False)
        label_sonuc.configure(text=f"✅ {fatura_tipi}.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_ve_isle_mustahsil(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Müstahsil XML klasörünü seçin")
    if secilen_klasor:
        xml_mustahsil_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac, "mustahsil_listesi")

# ----------------- Hepsiburada XML -----------------

def xml_hepsiburada_alis_oku_ve_yaz(klasor, label_sonuc, buton_ac):
    global SON_DOSYA_YOLU
    veriler = []
    ns = {
        "cbc": "urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2",
        "cac": "urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2",
    }
    
    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            
            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)
            
            # Alış faturası: Cari = Satıcı (Supplier)
            supplier_party = root.find(".//cac:AccountingSupplierParty/cac:Party", namespaces=ns)
            carikodu = supplier_party.findtext("cac:PartyIdentification/cbc:ID", namespaces=ns) if supplier_party is not None else ""
            cariadi = cari_adi_al(supplier_party, ns) if supplier_party is not None else "Bilinmeyen Cari"
            
            doviz = root.findtext(".//cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"
            
            for line in root.findall(".//cac:InvoiceLine", namespaces=ns):
                miktar = int(float(line.findtext("cbc:InvoicedQuantity", default="0", namespaces=ns)))
                fiyat = float(line.findtext("cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext("cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext("cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns) or ""
                
                # İskonto
                iskonto_element = line.find("cac:AllowanceCharge/cbc:Amount", namespaces=ns)
                iskonto = float(iskonto_element.text.replace(",", ".")) if iskonto_element is not None else 0.0
                
                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except:
                    kdv_orani = 0
                
                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Alış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": "ADET",
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto if iskonto else "",
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat * miktar,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")
    
    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, "hepsiburada_alis.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text="✅ hepsiburada_alis.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def xml_hepsiburada_satis_oku_ve_yaz(klasor, label_sonuc, buton_ac):
    global SON_DOSYA_YOLU
    veriler = []
    ns = {
        "cbc": "urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2",
        "cac": "urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2",
    }
    
    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            
            tarih_raw = root.findtext(".//cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext(".//cbc:ID", namespaces=ns)
            
            # Satış faturası: Cari = Alıcı (Customer)
            customer_party = root.find(".//cac:AccountingCustomerParty/cac:Party", namespaces=ns)
            carikodu = customer_party.findtext("cac:PartyIdentification/cbc:ID", namespaces=ns) if customer_party is not None else ""
            cariadi = cari_adi_al(customer_party, ns) if customer_party is not None else "Bilinmeyen Cari"
            
            doviz = root.findtext(".//cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"
            
            for line in root.findall(".//cac:InvoiceLine", namespaces=ns):
                miktar = int(float(line.findtext("cbc:InvoicedQuantity", default="0", namespaces=ns)))
                fiyat = float(line.findtext("cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext("cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext("cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns) or ""
                
                # İskonto
                iskonto_element = line.find("cac:AllowanceCharge/cbc:Amount", namespaces=ns)
                iskonto = float(iskonto_element.text.replace(",", ".")) if iskonto_element is not None else 0.0
                
                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except:
                    kdv_orani = 0
                
                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Satış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": "ADET",
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto if iskonto else "",
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat * miktar,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")
    
    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, "hepsiburada_satis.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text="✅ hepsiburada_satis.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_hepsiburada_alis(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Hepsiburada Alış Faturası klasörünü seçin")
    if secilen_klasor:
        xml_hepsiburada_alis_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac)

def klasor_sec_hepsiburada_satis(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Hepsiburada Satış Faturası klasörünü seçin")
    if secilen_klasor:
        xml_hepsiburada_satis_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac)

# ----------------- Trendyol XML -----------------

def xml_trendyol_alis_oku_ve_yaz(klasor, label_sonuc, buton_ac):
    # Henüz desteklenmiyor - örnek fatura geldiğinde eklenecek
    pass

def xml_trendyol_satis_oku_ve_yaz(klasor, label_sonuc, buton_ac):
    global SON_DOSYA_YOLU
    veriler = []
    ns = {
        "cbc": "urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2",
        "cac": "urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2",
    }
    
    for xml_dosya in glob.glob(os.path.join(klasor, "**", "*.xml"), recursive=True):
        try:
            tree = ET.parse(xml_dosya)
            root = tree.getroot()
            
            tarih_raw = root.findtext("cbc:IssueDate", namespaces=ns)
            tarih = datetime.strptime(tarih_raw, "%Y-%m-%d").strftime("%d.%m.%Y")
            evrakno = root.findtext("cbc:ID", namespaces=ns)
            
            # Satış faturası: Cari = Alıcı (Customer)
            customer_party = root.find("cac:AccountingCustomerParty/cac:Party", namespaces=ns)
            carikodu = customer_party.findtext("cac:PartyIdentification/cbc:ID", namespaces=ns) if customer_party is not None else ""
            cariadi = cari_adi_al(customer_party, ns) if customer_party is not None else "Bilinmeyen Cari"
            
            doviz = root.findtext("cbc:DocumentCurrencyCode", default="TRY", namespaces=ns)
            if doviz == "TRY": doviz = "TL"
            
            # Fatura düzeyinde toplam iskonto
            toplam_iskonto_element = root.findtext("cac:LegalMonetaryTotal/cbc:AllowanceTotalAmount", namespaces=ns)
            toplam_iskonto = float(toplam_iskonto_element.replace(",", ".")) if toplam_iskonto_element else 0.0
            
            lines = root.findall("cac:InvoiceLine", namespaces=ns)
            for idx, line in enumerate(lines):
                miktar = int(float(line.findtext("cbc:InvoicedQuantity", default="0", namespaces=ns)))
                fiyat = float(line.findtext("cac:Price/cbc:PriceAmount", default="0", namespaces=ns).replace(",", "."))
                stokadi = line.findtext("cac:Item/cbc:Name", namespaces=ns)
                stokkodu = line.findtext("cac:Item/cac:SellersItemIdentification/cbc:ID", namespaces=ns) or ""
                
                # İskonto: toplam iskontoyu sadece ilk satıra yaz, yoksa boş
                iskonto = toplam_iskonto if idx == 0 and toplam_iskonto else ""
                
                kdv_raw = line.findtext("cac:TaxTotal/cac:TaxSubtotal/cbc:Percent", namespaces=ns)
                try:
                    kdv_orani = int(float(kdv_raw))
                except:
                    kdv_orani = 0
                
                veriler.append({
                    "SIRANO": "",
                    "TARIH": tarih,
                    "EVRAKNO": evrakno,
                    "EVRAKTURU": "Satış Faturası",
                    "CARIKODU": carikodu,
                    "CARIADI": cariadi,
                    "STOKKODU": stokkodu,
                    "STOKADI": stokadi,
                    "BIRIM": "ADET",
                    "MIKTAR": miktar,
                    "DOVIZC": doviz,
                    "BIRIMFIYAT": fiyat,
                    "INDTL": iskonto,
                    "KDVY": kdv_orani,
                    "TUTARTL": fiyat * miktar,
                    "FATURACK": "",
                })
        except Exception as e:
            label_sonuc.configure(text=f"HATA: {os.path.basename(xml_dosya)} - {e}", text_color="red")
    
    if veriler:
        dosya_yolu = os.path.join(MASAUSTU, "trendyol_satis.xlsx")
        kaydet_excel(veriler, dosya_yolu)
        label_sonuc.configure(text="✅ trendyol_satis.xlsx masaüstüne kaydedildi.", text_color="green")
        buton_ac.configure(state="normal")
        SON_DOSYA_YOLU = dosya_yolu
    else:
        label_sonuc.configure(text="Geçerli XML dosyası bulunamadı.", text_color="red")

def klasor_sec_trendyol_alis(label_sonuc, buton_ac):
    # Henüz Trendyol Alış faturası desteği yok
    label_sonuc.configure(text="⚠️ Trendyol Alış faturası henüz desteklenmiyor.", text_color="orange")

def klasor_sec_trendyol_satis(label_sonuc, buton_ac):
    secilen_klasor = filedialog.askdirectory(title="Trendyol Satış Faturası klasörünü seçin")
    if secilen_klasor:
        xml_trendyol_satis_oku_ve_yaz(secilen_klasor, label_sonuc, buton_ac)

# ----------------- GUI -----------------

app = ctk.CTk()
app.title("Fatura Yönetim Sistemi")
app.geometry("650x610")
app.resizable(False, False)

# Açık popup menüyü takip etmek için
aktif_menu = None

# ----------------- Sayfa Yönetimi -----------------

def show_welcome():
    """Ana karşılama sayfasını göster"""
    xml_frame.pack_forget()
    kontrol_frame.pack_forget()
    welcome_frame.pack(expand=True, fill="both", padx=40, pady=30)

def show_xml():
    """XML aktarma sayfasını göster"""
    welcome_frame.pack_forget()
    kontrol_frame.pack_forget()
    xml_frame.pack(expand=True, fill="both", padx=40, pady=30)

def show_kontrol():
    """Kontrol sayfasını göster"""
    welcome_frame.pack_forget()
    xml_frame.pack_forget()
    kontrol_frame.pack(expand=True, fill="both", padx=40, pady=30)

# ----------------- Popup Menü Fonksiyonları -----------------

popup_aktif = False  # Popup durumunu takip et

def popup_menu_kapat():
    global aktif_menu, popup_aktif
    popup_aktif = False
    if aktif_menu is not None:
        try:
            aktif_menu.destroy()
        except Exception:
            # Pencere zaten kapatılmış olabilir
            pass
        aktif_menu = None

def on_app_click(event):
    """Ana pencereye tıklanınca popup'ı kapat"""
    global popup_aktif
    if popup_aktif and aktif_menu is not None:
        # Popup içine tıklanmadıysa kapat
        widget = event.widget
        # Popup'ın child'ı değilse kapat
        if str(widget).startswith(".!ctktoplevel"):
            return  # Popup içi, kapatma
        popup_menu_kapat()

def show_popup_menu(button, label_sonuc, buton_ac, alis_func, satis_func):
    """Genel popup menü fonksiyonu - Alış/Satış seçenekleri gösterir"""
    global aktif_menu, popup_aktif
    popup_menu_kapat()
    
    x = button.winfo_rootx()
    y = button.winfo_rooty() + button.winfo_height() + 5
    
    aktif_menu = ctk.CTkToplevel(app)
    aktif_menu.geometry(f"160x90+{x}+{y}")
    aktif_menu.overrideredirect(True)
    aktif_menu.attributes("-topmost", True)
    aktif_menu.configure(fg_color="#3b3b3b")
    
    ctk.CTkButton(aktif_menu, text="📥 Alış Faturası", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), alis_func(label_sonuc, buton_ac)]).pack(pady=(8, 4), padx=5)
    ctk.CTkButton(aktif_menu, text="📤 Satış Faturası", width=150, height=35,
                  command=lambda: [popup_menu_kapat(), satis_func(label_sonuc, buton_ac)]).pack(pady=(0, 8), padx=5)
    
    # Popup aktif olarak işaretle (biraz gecikmeyle)
    def activate_popup():
        global popup_aktif
        popup_aktif = True
    app.after(150, activate_popup)
    
    # Escape tuşu ile kapatma
    aktif_menu.bind("<Escape>", lambda e: popup_menu_kapat())

def show_hepsiburada_menu(button, label_sonuc, buton_ac):
    show_popup_menu(button, label_sonuc, buton_ac, klasor_sec_hepsiburada_alis, klasor_sec_hepsiburada_satis)

def show_trendyol_menu(button, label_sonuc, buton_ac):
    show_popup_menu(button, label_sonuc, buton_ac, klasor_sec_trendyol_alis, klasor_sec_trendyol_satis)

# ================= KARŞILAMA SAYFASI =================

welcome_frame = ctk.CTkFrame(app, fg_color="transparent")

welcome_inner = ctk.CTkFrame(welcome_frame, fg_color="#2b2b2b", corner_radius=15)
welcome_inner.pack(expand=True, fill="both")

# Başlık
welcome_title = ctk.CTkLabel(
    welcome_inner, 
    text="📄 Fatura Yönetim Sistemi", 
    font=ctk.CTkFont(size=28, weight="bold")
)
welcome_title.pack(pady=(30, 8))

welcome_desc = ctk.CTkLabel(
    welcome_inner,
    text="Yapmak istediğiniz işlemi seçin",
    font=ctk.CTkFont(size=16),
    text_color="#a0a0a0"
)
welcome_desc.pack(pady=(0, 20))

# Kart Container - tüm alanı kaplasın
cards_frame = ctk.CTkFrame(welcome_inner, fg_color="transparent")
cards_frame.pack(expand=True, fill="both", padx=25, pady=(0, 25))
cards_frame.grid_columnconfigure(0, weight=1)
cards_frame.grid_columnconfigure(1, weight=1)
cards_frame.grid_rowconfigure(0, weight=1)

# Kart hover efekti ve tıklama için yardımcı fonksiyon
def make_card_clickable(card, all_widgets, command):
    """Karta hover efekti ve tıklama özelliği ekler"""
    normal_color = "#363636"
    hover_color = "#454545"
    click_color = "#505050"
    
    def on_enter(e):
        card.configure(fg_color=hover_color)
    
    def on_leave(e):
        card.configure(fg_color=normal_color)
    
    def on_click(e):
        card.configure(fg_color=click_color)
        card.after(100, lambda: [card.configure(fg_color=normal_color), command()])
    
    # Kart ve tüm child widget'lara event bağla
    for widget in all_widgets:
        widget.bind("<Enter>", on_enter)
        widget.bind("<Leave>", on_leave)
        widget.bind("<Button-1>", on_click)
        widget.configure(cursor="hand2")

# XML Kartı
xml_card = ctk.CTkFrame(cards_frame, fg_color="#363636", corner_radius=12)
xml_card.grid(row=0, column=0, padx=(0, 10), sticky="nsew")

xml_card_content = ctk.CTkFrame(xml_card, fg_color="transparent")
xml_card_content.pack(expand=True)

xml_icon = ctk.CTkLabel(xml_card_content, text="📁", font=ctk.CTkFont(size=56))
xml_icon.pack(pady=(0, 15))

xml_card_title = ctk.CTkLabel(xml_card_content, text="XML Aktarma", font=ctk.CTkFont(size=22, weight="bold"))
xml_card_title.pack(pady=(0, 10))

xml_card_desc = ctk.CTkLabel(
    xml_card_content, 
    text="XML faturalarını\nExcel'e aktarın",
    font=ctk.CTkFont(size=14),
    text_color="#b0b0b0",
    justify="center"
)
xml_card_desc.pack()

# XML kartını tıklanabilir yap
make_card_clickable(xml_card, [xml_card, xml_card_content, xml_icon, xml_card_title, xml_card_desc], show_xml)

# Kontrol Kartı
kontrol_card = ctk.CTkFrame(cards_frame, fg_color="#363636", corner_radius=12)
kontrol_card.grid(row=0, column=1, padx=(10, 0), sticky="nsew")

kontrol_card_content = ctk.CTkFrame(kontrol_card, fg_color="transparent")
kontrol_card_content.pack(expand=True)

kontrol_icon = ctk.CTkLabel(kontrol_card_content, text="🔍", font=ctk.CTkFont(size=56))
kontrol_icon.pack(pady=(0, 15))

kontrol_card_title = ctk.CTkLabel(kontrol_card_content, text="Fatura Kontrol", font=ctk.CTkFont(size=22, weight="bold"))
kontrol_card_title.pack(pady=(0, 10))

kontrol_card_desc = ctk.CTkLabel(
    kontrol_card_content, 
    text="Zirve ve portal\nfaturalarını karşılaştırın",
    font=ctk.CTkFont(size=14),
    text_color="#b0b0b0",
    justify="center"
)
kontrol_card_desc.pack()

# Kontrol kartını tıklanabilir yap
make_card_clickable(kontrol_card, [kontrol_card, kontrol_card_content, kontrol_icon, kontrol_card_title, kontrol_card_desc], show_kontrol)

# ================= XML SAYFASI =================

xml_frame = ctk.CTkFrame(app, fg_color="transparent")

xml_inner = ctk.CTkFrame(xml_frame, fg_color="#2b2b2b", corner_radius=15)
xml_inner.pack(expand=True, fill="both")

# Geri Butonu
xml_back_frame = ctk.CTkFrame(xml_inner, fg_color="transparent")
xml_back_frame.pack(fill="x", padx=15, pady=(15, 0))

xml_back_btn = ctk.CTkButton(
    xml_back_frame, 
    text="← Geri", 
    width=80, 
    height=30,
    fg_color="transparent",
    border_width=1,
    border_color="#555555",
    hover_color="#404040",
    command=show_welcome
)
xml_back_btn.pack(side="left")

# Başlık
xml_title = ctk.CTkLabel(xml_inner, text="📄 XML Aktarma", font=ctk.CTkFont(size=24, weight="bold"))
xml_title.pack(pady=(10, 5))

xml_aciklama = ctk.CTkLabel(
    xml_inner,
    text="XML formatındaki faturaları seçerek Excel çıktısı oluşturabilirsiniz.",
    font=ctk.CTkFont(size=14),
    text_color="#d0d0d0",
    justify="center"
)
xml_aciklama.pack(pady=(0, 20))

# Logo ve Uyumsoft Butonları
xml_buton_frame = ctk.CTkFrame(xml_inner, fg_color="transparent")
xml_buton_frame.pack(pady=(0, 10))

# Logo XML
logo_col = ctk.CTkFrame(xml_buton_frame, fg_color="transparent")
logo_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(logo_col, text="Logo XML", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(logo_col, text="📂 Aktar", width=160, height=40,
              command=lambda: klasor_sec_ve_isle_logo(xml_label_sonuc, xml_buton_ac)).pack()

# Uyumsoft XML
uyumsoft_col = ctk.CTkFrame(xml_buton_frame, fg_color="transparent")
uyumsoft_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(uyumsoft_col, text="Uyumsoft XML", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(uyumsoft_col, text="🧾 Aktar", width=160, height=40,
              command=lambda: klasor_sec_ve_isle_uyumsoft(xml_label_sonuc, xml_buton_ac)).pack()

# Müstahsil XML
xml_mustahsil_frame = ctk.CTkFrame(xml_inner, fg_color="transparent")
xml_mustahsil_frame.pack(pady=(10, 5))
ctk.CTkLabel(xml_mustahsil_frame, text="Müstahsil XML", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(xml_mustahsil_frame, text="🥔 Aktar", width=160, height=40,
              command=lambda: klasor_sec_ve_isle_mustahsil(xml_label_sonuc, xml_buton_ac)).pack()

# Hepsiburada ve Trendyol butonları
xml_pazar_yeri_frame = ctk.CTkFrame(xml_inner, fg_color="transparent")
xml_pazar_yeri_frame.pack(pady=(15, 5))

# Hepsiburada
hb_col = ctk.CTkFrame(xml_pazar_yeri_frame, fg_color="transparent")
hb_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(hb_col, text="Hepsiburada", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
hb_buton = ctk.CTkButton(hb_col, text="🛒 Aktar", width=160, height=40)
hb_buton.pack()

# Trendyol
ty_col = ctk.CTkFrame(xml_pazar_yeri_frame, fg_color="transparent")
ty_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(ty_col, text="Trendyol", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ty_buton = ctk.CTkButton(ty_col, text="🛍️ Aktar", width=160, height=40)
ty_buton.pack()

# Sonuç mesajı
xml_label_sonuc = ctk.CTkLabel(xml_inner, text="", font=ctk.CTkFont(size=16))
xml_label_sonuc.pack(pady=10)

# Excel aç butonu
xml_buton_ac = ctk.CTkButton(xml_inner, text="📂 Oluşturulan Excel Dosyasını Aç", command=dosyayi_ac, state="disabled", width=280, height=40)
xml_buton_ac.pack(pady=(10, 30))

# Popup menü komutlarını bağla (label_sonuc ve buton_ac tanımlandıktan sonra)
hb_buton.configure(command=lambda: show_hepsiburada_menu(hb_buton, xml_label_sonuc, xml_buton_ac))
ty_buton.configure(command=lambda: show_trendyol_menu(ty_buton, xml_label_sonuc, xml_buton_ac))

# ================= KONTROL SAYFASI =================

kontrol_frame = ctk.CTkFrame(app, fg_color="transparent")

kontrol_inner = ctk.CTkFrame(kontrol_frame, fg_color="#2b2b2b", corner_radius=15)
kontrol_inner.pack(expand=True, fill="both")

# Geri Butonu
kontrol_back_frame = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_back_frame.pack(fill="x", padx=15, pady=(15, 0))

kontrol_back_btn = ctk.CTkButton(
    kontrol_back_frame, 
    text="← Geri", 
    width=80, 
    height=30,
    fg_color="transparent",
    border_width=1,
    border_color="#555555",
    hover_color="#404040",
    command=show_welcome
)
kontrol_back_btn.pack(side="left")

# Başlık
kontrol_title = ctk.CTkLabel(kontrol_inner, text="🔍 Fatura Kontrol", font=ctk.CTkFont(size=24, weight="bold"))
kontrol_title.pack(pady=(10, 5))

kontrol_aciklama = ctk.CTkLabel(
    kontrol_inner,
    text="Zirve ve portal faturalarını karşılaştırarak\neksik veya hatalı girişleri tespit edin.",
    font=ctk.CTkFont(size=14),
    text_color="#d0d0d0",
    justify="center"
)
kontrol_aciklama.pack(pady=(0, 20))

# Kontrol popup fonksiyonu
def show_kontrol_popup(portal_adi, kontrol_func=None):
    """Excel karşılaştırma popup'ı göster"""
    popup = ctk.CTkToplevel(app)
    popup.title(f"{portal_adi} Kontrol")
    popup.geometry("500x360")
    popup.resizable(False, False)
    popup.attributes("-topmost", True)
    popup.grab_set()  # Modal yap
    
    # Pencereyi ortala
    popup.update_idletasks()
    x = app.winfo_x() + (app.winfo_width() // 2) - 250
    y = app.winfo_y() + (app.winfo_height() // 2) - 180
    popup.geometry(f"500x360+{x}+{y}")
    
    # Seçilen dosyaları takip et
    secilen_dosyalar = {"zirve": None, "portal": []}
    
    # Başlık
    ctk.CTkLabel(popup, text=f"🔍 {portal_adi} Fatura Kontrol", 
                 font=ctk.CTkFont(size=20, weight="bold")).pack(pady=(20, 5))
    ctk.CTkLabel(popup, text="Karşılaştırmak için Excel dosyası/dosyaları seçin", 
                 font=ctk.CTkFont(size=12), text_color="#a0a0a0").pack(pady=(0, 20))
    
    # Dosya seçim alanları
    files_frame = ctk.CTkFrame(popup, fg_color="transparent")
    files_frame.pack(fill="x", padx=30)
    
    # Sol: Zirve Excel
    zirve_frame = ctk.CTkFrame(files_frame, fg_color="#363636", corner_radius=10)
    zirve_frame.grid(row=0, column=0, padx=(0, 10), sticky="nsew")
    files_frame.grid_columnconfigure(0, weight=1)
    
    ctk.CTkLabel(zirve_frame, text="Zirve Excel", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 5))
    zirve_status = ctk.CTkLabel(zirve_frame, text="📄 Dosya seçilmedi", font=ctk.CTkFont(size=11), text_color="#808080")
    zirve_status.pack(pady=(0, 10))
    
    def select_zirve():
        # Popup'ı gizle
        popup.withdraw()
        popup.update_idletasks()
        
        dosya = filedialog.askopenfilename(
            title="Zirve Excel Dosyası Seçin",
            filetypes=[("Excel Dosyaları", "*.xlsx *.xls")]
        )
        
        # Popup'ı tekrar göster
        popup.deiconify()
        popup.lift()
        popup.focus_force()
        
        if dosya:
            secilen_dosyalar["zirve"] = dosya
            dosya_adi = os.path.basename(dosya)
            if len(dosya_adi) > 20:
                dosya_adi = dosya_adi[:17] + "..."
            zirve_status.configure(text=f"✅ {dosya_adi}", text_color="#4CAF50")
            check_ready()
    
    ctk.CTkButton(zirve_frame, text="📂 Seç", width=100, height=32, command=select_zirve).pack(pady=(0, 15))
    
    # Sağ: Portal Excel (Çoklu Seçim)
    portal_frame = ctk.CTkFrame(files_frame, fg_color="#363636", corner_radius=10)
    portal_frame.grid(row=0, column=1, padx=(10, 0), sticky="nsew")
    files_frame.grid_columnconfigure(1, weight=1)
    
    ctk.CTkLabel(portal_frame, text=f"{portal_adi} Excel", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 5))
    portal_status = ctk.CTkLabel(portal_frame, text="📄 Dosya seçilmedi", font=ctk.CTkFont(size=11), text_color="#808080")
    portal_status.pack(pady=(0, 10))
    
    def select_portal():
        # Popup'ı gizle
        popup.withdraw()
        popup.update_idletasks()
        
        dosyalar = filedialog.askopenfilenames(
            title=f"{portal_adi} Excel Dosyası(ları) Seçin (Çoklu seçim için Ctrl tuşu)",
            filetypes=[("Excel Dosyaları", "*.xlsx *.xls")]
        )
        
        # Popup'ı tekrar göster
        popup.deiconify()
        popup.lift()
        popup.focus_force()
        
        if dosyalar:
            secilen_dosyalar["portal"] = list(dosyalar)
            dosya_sayisi = len(dosyalar)
            
            if dosya_sayisi == 1:
                dosya_adi = os.path.basename(dosyalar[0])
                if len(dosya_adi) > 18:
                    dosya_adi = dosya_adi[:15] + "..."
                portal_status.configure(text=f"✅ {dosya_adi}", text_color="#4CAF50")
            else:
                portal_status.configure(text=f"✅ {dosya_sayisi} dosya seçildi", text_color="#4CAF50")
            
            check_ready()
    
    ctk.CTkButton(portal_frame, text="📂 Dosyalar Seç", width=110, height=32, command=select_portal).pack(pady=(0, 15))
    
    # Kontrol Et butonu (başlangıçta devre dışı)
    def run_kontrol():
        if secilen_dosyalar["zirve"] and secilen_dosyalar["portal"]:
            popup.destroy()
            if kontrol_func:
                kontrol_func(secilen_dosyalar["zirve"], secilen_dosyalar["portal"])
            else:
                # Henüz implementasyon yok
                show_info_popup("Bu özellik henüz aktif değil.")
    
    kontrol_btn = ctk.CTkButton(popup, text="🔍 Kontrol Et", width=200, height=40, 
                                 command=run_kontrol, state="disabled",
                                 fg_color="#555555", hover_color="#555555")
    kontrol_btn.pack(pady=(25, 15))
    
    def check_ready():
        # Zirve seçilmiş ve Portal'de en az 1 dosya seçilmişse
        if secilen_dosyalar["zirve"] and len(secilen_dosyalar["portal"]) > 0:
            kontrol_btn.configure(state="normal", fg_color="#1f6aa5", hover_color="#144870")
    
    # İptal butonu
    ctk.CTkButton(popup, text="İptal", width=100, height=30, 
                  fg_color="transparent", border_width=1, border_color="#555555",
                  hover_color="#404040", command=popup.destroy).pack()

def show_info_popup(mesaj):
    """Bilgi popup'ı göster"""
    popup = ctk.CTkToplevel(app)
    popup.title("Bilgi")
    popup.geometry("300x150")
    popup.resizable(False, False)
    popup.attributes("-topmost", True)
    popup.grab_set()
    
    popup.update_idletasks()
    x = app.winfo_x() + (app.winfo_width() // 2) - 150
    y = app.winfo_y() + (app.winfo_height() // 2) - 75
    popup.geometry(f"300x150+{x}+{y}")
    
    ctk.CTkLabel(popup, text="⚠️", font=ctk.CTkFont(size=36)).pack(pady=(20, 5))
    ctk.CTkLabel(popup, text=mesaj, font=ctk.CTkFont(size=14)).pack(pady=(0, 15))
    ctk.CTkButton(popup, text="Tamam", width=100, height=32, command=popup.destroy).pack()

# Kontrol fonksiyonları - her portal için
def kontrol_logo():
    show_info_popup("Bu özellik henüz aktif değil.")

def uyumsoft_kontrol_callback(zirve_path, portal_path):
    """Uyumsoft karşılaştırma sonrası çağrılır"""
    sonuc_dosyasi, mesaj = uyumsoft_karsilastir(zirve_path, portal_path)
    
    if sonuc_dosyasi:
        # Başarılı - sonuç popup'ı göster
        popup = ctk.CTkToplevel(app)
        popup.title("Kontrol Sonucu")
        popup.geometry("420x320")
        popup.resizable(False, False)
        popup.attributes("-topmost", True)
        popup.grab_set()
        
        # Pencereyi ortala
        popup.update_idletasks()
        x = app.winfo_x() + (app.winfo_width() // 2) - 210
        y = app.winfo_y() + (app.winfo_height() // 2) - 160
        popup.geometry(f"420x320+{x}+{y}")
        
        # İçerik
        ctk.CTkLabel(popup, text="🎉", font=ctk.CTkFont(size=36)).pack(pady=(20, 5))
        ctk.CTkLabel(popup, text=mesaj, font=ctk.CTkFont(size=13), 
                     justify="left").pack(pady=(0, 20), padx=20)
        
        # Butonlar
        btn_frame = ctk.CTkFrame(popup, fg_color="transparent")
        btn_frame.pack(pady=(0, 20))
        
        ctk.CTkButton(btn_frame, text="📂 Dosyayı Aç", width=150, height=35, 
                      command=lambda: [popup.destroy(), dosyayi_ac()]).pack(side="left", padx=5)
        ctk.CTkButton(btn_frame, text="Kapat", width=100, height=35, 
                      fg_color="transparent", border_width=1, border_color="#555555",
                      command=popup.destroy).pack(side="left", padx=5)
    else:
        # Hata - hata popup'ı göster
        show_info_popup(mesaj)

def kontrol_uyumsoft():
    show_kontrol_popup("Uyumsoft", uyumsoft_kontrol_callback)  # Callback ile

def vega_kontrol_callback(zirve_path, portal_path):
    """Vega karşılaştırma sonrası çağrılır"""
    sonuc_dosyasi, mesaj = vega_karsilastir(zirve_path, portal_path)
    
    if sonuc_dosyasi:
        # Başarılı - sonuç popup'ı göster
        popup = ctk.CTkToplevel(app)
        popup.title("Kontrol Sonucu")
        popup.geometry("420x320")
        popup.resizable(False, False)
        popup.attributes("-topmost", True)
        popup.grab_set()
        
        # Pencereyi ortala
        popup.update_idletasks()
        x = app.winfo_x() + (app.winfo_width() // 2) - 210
        y = app.winfo_y() + (app.winfo_height() // 2) - 160
        popup.geometry(f"420x320+{x}+{y}")
        
        # İçerik
        ctk.CTkLabel(popup, text="🎉", font=ctk.CTkFont(size=36)).pack(pady=(20, 5))
        ctk.CTkLabel(popup, text=mesaj, font=ctk.CTkFont(size=13), 
                     justify="left").pack(pady=(0, 20), padx=20)
        
        # Butonlar
        btn_frame = ctk.CTkFrame(popup, fg_color="transparent")
        btn_frame.pack(pady=(0, 20))
        
        ctk.CTkButton(btn_frame, text="📂 Dosyayı Aç", width=150, height=35, 
                      command=lambda: [popup.destroy(), dosyayi_ac()]).pack(side="left", padx=5)
        ctk.CTkButton(btn_frame, text="Kapat", width=100, height=35, 
                      fg_color="transparent", border_width=1, border_color="#555555",
                      command=popup.destroy).pack(side="left", padx=5)
    else:
        # Hata - hata popup'ı göster
        show_info_popup(mesaj)

def kontrol_hizlibilisim():
    show_info_popup("Bu özellik henüz aktif değil.")

def kontrol_izibiz():
    show_info_popup("Bu özellik henüz aktif değil.")

def kontrol_vega():
    show_kontrol_popup("Vega", vega_kontrol_callback)  # Callback ile

def kontrol_qnb():
    show_info_popup("Bu özellik henüz aktif değil.")

def kontrol_hepsiburada():
    show_info_popup("Bu özellik henüz aktif değil.")

def kontrol_trendyol():
    show_info_popup("Bu özellik henüz aktif değil.")

# Satır 1: Logo ve Uyumsoft
kontrol_row1 = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_row1.pack(pady=(0, 10))

logo_kontrol_col = ctk.CTkFrame(kontrol_row1, fg_color="transparent")
logo_kontrol_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(logo_kontrol_col, text="Logo", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(logo_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_logo).pack()

uyumsoft_kontrol_col = ctk.CTkFrame(kontrol_row1, fg_color="transparent")
uyumsoft_kontrol_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(uyumsoft_kontrol_col, text="Uyumsoft", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(uyumsoft_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_uyumsoft).pack()

# Satır 2: Hızlıbilişim ve İzibiz
kontrol_row2 = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_row2.pack(pady=(10, 10))

hizli_kontrol_col = ctk.CTkFrame(kontrol_row2, fg_color="transparent")
hizli_kontrol_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(hizli_kontrol_col, text="Hızlıbilişim", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(hizli_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_hizlibilisim).pack()

izibiz_kontrol_col = ctk.CTkFrame(kontrol_row2, fg_color="transparent")
izibiz_kontrol_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(izibiz_kontrol_col, text="İzibiz", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(izibiz_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_izibiz).pack()

# Satır 3: Vega ve QNB
kontrol_row3 = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_row3.pack(pady=(10, 10))

vega_kontrol_col = ctk.CTkFrame(kontrol_row3, fg_color="transparent")
vega_kontrol_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(vega_kontrol_col, text="Vega", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(vega_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_vega).pack()

qnb_kontrol_col = ctk.CTkFrame(kontrol_row3, fg_color="transparent")
qnb_kontrol_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(qnb_kontrol_col, text="QNB", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(qnb_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_qnb).pack()

# Satır 4: Hepsiburada ve Trendyol
kontrol_row4 = ctk.CTkFrame(kontrol_inner, fg_color="transparent")
kontrol_row4.pack(pady=(10, 10))

hb_kontrol_col = ctk.CTkFrame(kontrol_row4, fg_color="transparent")
hb_kontrol_col.grid(row=0, column=0, padx=40)
ctk.CTkLabel(hb_kontrol_col, text="Hepsiburada", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(hb_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_hepsiburada).pack()

ty_kontrol_col = ctk.CTkFrame(kontrol_row4, fg_color="transparent")
ty_kontrol_col.grid(row=0, column=1, padx=40)
ctk.CTkLabel(ty_kontrol_col, text="Trendyol", font=ctk.CTkFont(size=18)).pack(pady=(0, 10))
ctk.CTkButton(ty_kontrol_col, text="🔍 Kontrol", width=160, height=40,
              command=kontrol_trendyol).pack()

# Sonuç mesajı
kontrol_label_sonuc = ctk.CTkLabel(kontrol_inner, text="", font=ctk.CTkFont(size=16))
kontrol_label_sonuc.pack(pady=10)

# ================= BAŞLANGIÇ =================

# Ana pencereye tıklanınca popup menüyü kapat
app.bind_all("<Button-1>", on_app_click)

# Karşılama sayfasını göster
show_welcome()

app.mainloop()